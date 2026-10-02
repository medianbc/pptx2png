# ==========================================
# yandex_flow_core.py — ЯДРО (v4.0, часть 1/3)
# ==========================================
# Утилиты, категории, защита in-flight операций, workers,
# промпты (рендер), watchdog, cleanup, конвертация.
#
# НЕ содержит хендлеров @router.* — они в yandex_flow_handlers.py.
# НЕ импортирует yandex_flow_handlers, чтобы не создать цикл.
# ==========================================

import asyncio
import html as html_module
import logging
import os
import secrets
import shutil
import tempfile
import time
import urllib.parse
from pathlib import Path
from typing import Optional, Callable, Awaitable, List, Dict, Tuple

from aiogram import types, Bot
from aiogram.exceptions import TelegramBadRequest
from aiogram.types import InlineKeyboardButton
from aiogram.utils.keyboard import InlineKeyboardBuilder

import yandex_state
from yandex_state import (
    sessions,
    yd_session_lock,
    yd_active_tasks,
    yd_release,
)

from yandex_disk import YandexDiskError
from structure import safe_folder_name
from sermon_detector import find_sermon_range
from utils import extract_speaker_notes
from converter_engine import (
    convert_all_pngs,
    create_zip_stream,
    ppt_to_pptx_crossplatform,
    librenormalize_to_pptx,
)


# ==========================================
# МЕТАДАННЫЕ КАТЕГОРИЙ
# ==========================================

_DEFAULT_CATEGORY_ORDER = ["opening", "prayer", "sermon"]

_DEFAULT_CATEGORY_META = {
    "opening": {
        "name": "Начало",
        "emoji": "🌅",
        "folder_attr": "opening_folder",
        "zip_suffix": "начало",
        "settings_flag": "process_opening",
    },
    "prayer": {
        "name": "Молитва",
        "emoji": "🙏",
        "folder_attr": "prayer_folder",
        "zip_suffix": "молитва",
        "settings_flag": "process_prayer",
    },
    "sermon": {
        "name": "Проповедь",
        "emoji": "🎯",
        "folder_attr": "sermon_folder",
        "zip_suffix": "проповедь",
        "settings_flag": "process_sermon",
    },
}


def _cat_order() -> List[str]:
    """Возвращает актуальный порядок категорий."""
    return getattr(
        yandex_state.config, "category_order", _DEFAULT_CATEGORY_ORDER
    ) or _DEFAULT_CATEGORY_ORDER


def _cat_meta() -> Dict[str, dict]:
    """Возвращает актуальные метаданные категорий."""
    return getattr(
        yandex_state.config, "category_meta", _DEFAULT_CATEGORY_META
    ) or _DEFAULT_CATEGORY_META


def _cat_folder_path(cat: str, target_base: str) -> str:
    """Возвращает полный путь папки категории под target_base."""
    meta = _cat_meta().get(cat, {})
    folder_attr = meta.get("folder_attr")
    folder_name = (
        getattr(yandex_state.config, folder_attr, "")
        if folder_attr else ""
    )
    if not folder_name:
        folder_name = meta.get("name", cat)
    return f"{target_base}/{folder_name}"


# ==========================================
# БАЗОВЫЕ УТИЛИТЫ
# ==========================================

def _touch_task(task_dir: Path):
    """Обновляет mtime папки — чтобы cleaner не удалял активную задачу."""
    if task_dir and task_dir.exists():
        try:
            os.utime(task_dir, None)
        except Exception as e:
            logging.error(f"Ошибка touch для {task_dir}: {e}")


def _safe_delete_task_dir(task_dir: Path):
    """Безопасно удаляет папку задачи."""
    if task_dir and task_dir.exists():
        try:
            shutil.rmtree(task_dir)
            logging.info(f"🧹 Удалена папка задачи: {task_dir}")
        except Exception as e:
            logging.error(f"Ошибка удаления папки {task_dir}: {e}")


def _safe_unlink(path: Path):
    """Безопасно удаляет файл, логирует ошибку при неудаче."""
    if path is None:
        return
    try:
        if path.exists():
            path.unlink()
    except Exception as e:
        logging.warning(f"Не удалось удалить {path}: {e}")


async def _safe_edit(msg, text: str, **kwargs) -> bool:
    """Безопасный edit_text. Различает служебные TelegramBadRequest."""
    if msg is None:
        return False
    try:
        await msg.edit_text(text, **kwargs)
        return True
    except TelegramBadRequest as e:
        s = str(e)
        if "message is not modified" in s:
            return False
        if "query is too old" in s or "query ID is invalid" in s:
            logging.debug(f"[YD-EDIT] query устарел: {e}")
            return False
        logging.warning(f"TelegramBadRequest в edit_text: {e}")
        return False
    except Exception as e:
        logging.warning(
            f"Не удалось обновить статус-сообщение: {e}", exc_info=True
        )
        return False


async def _safe_answer(
    callback: types.CallbackQuery,
    text: str = "",
    show_alert: bool = False,
) -> None:
    """Безопасный callback.answer — глотает 'query is too old'."""
    try:
        await callback.answer(text, show_alert=show_alert)
    except TelegramBadRequest as e:
        s = str(e)
        if "query is too old" in s or "query ID is invalid" in s:
            logging.info(f"[YD-ANSWER] callback устарел: {e}")
        else:
            logging.warning(f"[YD-ANSWER] TelegramBadRequest: {e}")
    except Exception as e:
        logging.warning(f"[YD-ANSWER] не удалось ответить на callback: {e}")


def _format_ranges_text(ranges, start=None, end=None) -> str:
    """Форматирует список диапазонов в '1–2, 8–9'."""
    if ranges:
        return ", ".join(
            f"{s}–{e}" if s != e else str(s) for s, e in ranges
        )
    if start is not None and end is not None:
        if start != end:
            return f"{start}–{end}"
        return str(start)
    return "—"


def _ranges_to_start_end(ranges):
    """Возвращает (start, end) как min/max объединения ranges."""
    if not ranges:
        return None, None
    try:
        start = min(s for s, _ in ranges)
        end = max(e for _, e in ranges)
    except Exception:
        return None, None
    return start, end


def _normalize_item_ranges(item: dict) -> None:
    """Синхронизирует плоские ranges/start/end. Категории не трогает."""
    ranges = item.get("ranges")
    start = item.get("start")
    end = item.get("end")

    if ranges:
        s2, e2 = _ranges_to_start_end(ranges)
        item["start"] = s2
        item["end"] = e2
        return

    if start is not None and end is not None and start <= end:
        item["ranges"] = [(start, end)]
        return

    item["start"] = None
    item["end"] = None
    item["ranges"] = None


def _is_sermon_slide(item: dict, slide_idx: int) -> bool:
    """Совместимость со старым кодом. Проверяет плоские ranges."""
    ranges = item.get("ranges")
    if ranges:
        return any(s <= slide_idx <= e for s, e in ranges)
    start = item.get("start")
    end = item.get("end")
    return start is not None and start <= slide_idx <= end


def _yd_public_url(disk_path: str) -> str:
    """Строит корректный кликабельный URL Яндекс.Диска."""
    path = disk_path.lstrip("/")
    encoded = urllib.parse.quote(path, safe="/")
    return f"https://disk.yandex.ru/client/disk/{encoded}"


def _format_size(num_bytes: int) -> str:
    """Человекочитаемый размер файла."""
    if num_bytes < 1024:
        return f"{num_bytes} Б"
    if num_bytes < 1024 * 1024:
        return f"{num_bytes / 1024:.1f} КБ"
    if num_bytes < 1024 * 1024 * 1024:
        return f"{num_bytes / (1024 * 1024):.1f} МБ"
    return f"{num_bytes / (1024 * 1024 * 1024):.2f} ГБ"


def _get_cancel_keyboard(task_id: str) -> InlineKeyboardBuilder:
    """Клавиатура с единственной кнопкой «Отменить задачу»."""
    kb = InlineKeyboardBuilder()
    kb.row(
        InlineKeyboardButton(
            text="❌ Отменить задачу",
            callback_data=f"yd_task_cancel:{task_id}",
        ),
    )
    return kb


def _count_slides_for_range(start, end, total):
    """Возвращает количество слайдов в диапазоне [start, end]."""
    if start is None or end is None:
        return 0
    return max(0, min(end, total) - max(start, 1) + 1)


def _count_slides_in_ranges(ranges, total_slides: int) -> int:
    """Считает количество УНИКАЛЬНЫХ слайдов в объединении диапазонов."""
    if not ranges or total_slides <= 0:
        return 0

    unique_slides = set()
    for s, e in ranges:
        lo = max(1, s)
        hi = min(total_slides, e)
        if lo > hi:
            continue
        unique_slides.update(range(lo, hi + 1))
    return len(unique_slides)


# ==========================================
# PICKER STATUS (для хендлеров)
# ==========================================

def _picker_is_active(picker: Optional[dict]) -> bool:
    """Пикер активен, если статус != idle."""
    return _picker_status(picker) != "idle"


def _picker_status(picker: Optional[dict]) -> str:
    """
    Возвращает состояние пикера:
      • "none"    — сессии нет;
      • "draft"   — cmd_sunday в окне YD API запросов;
      • "working" — есть активная задача;
      • "idle"    — сессия готова, ждём выбора пользователя.
    """
    if not picker:
        return "none"
    if picker.get("_draft"):
        return "draft"
    if picker.get("processing"):
        return "working"
    for tid in picker.get("task_ids", []):
        task_sess = sessions.get(tid)
        if task_sess is not None and not task_sess.get("cancelled"):
            return "working"
    return "idle"


def _picker_has_live_tasks_locked(
    picker_key: str,
    exclude_task_id: Optional[str] = None,
) -> bool:
    """Проверяет наличие живых sibling-задач. Вызывать под yd_session_lock."""
    picker = sessions.get(picker_key)
    if picker is None:
        return False
    for tid in picker.get("task_ids", []):
        if tid == exclude_task_id:
            continue
        task_sess = sessions.get(tid)
        if task_sess is not None and not task_sess.get("cancelled"):
            return True
    return False


# ==========================================
# ДЕТЕКТ КАТЕГОРИЙ
# ==========================================

def _detect_all_categories(
    notes,
    notes_ok: bool,
    incomplete: bool,
) -> Dict[str, dict]:
    """Детектит все категории в заметках."""
    result: Dict[str, dict] = {}

    if not notes_ok or incomplete or not notes:
        for cat in _cat_order():
            result[cat] = {"ranges": None, "matches": [], "found": False}
        return result

    kw_map = {
        "opening": getattr(yandex_state.config, "opening_keywords", []),
        "prayer":  getattr(yandex_state.config, "prayer_keywords", []),
        "sermon":  getattr(yandex_state.config, "sermon_keywords", []),
    }

    for cat in _cat_order():
        kws = kw_map.get(cat) or []
        if not kws:
            result[cat] = {"ranges": None, "matches": [], "found": False}
            continue

        try:
            start, end, matches = find_sermon_range(notes, kws)
        except Exception as e:
            logging.error(
                f"[YD-CAT] find_sermon_range({cat}) упал: {e}",
                exc_info=True,
            )
            result[cat] = {"ranges": None, "matches": [], "found": False}
            continue

        if matches and len(matches) == 1:
            start, end = None, None

        ranges = [(start, end)] if start is not None else None
        result[cat] = {
            "ranges": ranges,
            "matches": list(matches) if matches else [],
            "found": bool(ranges or matches),
        }

    logging.debug(
        f"[YD-CAT] detect: "
        + ", ".join(
            f"{cat}={'ok' if result[cat]['found'] else 'no'}"
            for cat in _cat_order()
        )
    )
    return result


def _classify_slide(item: dict, slide_idx: int) -> Optional[str]:
    """
    Определяет, к какой категории относится слайд.
    Приоритет = порядок _cat_order() (opening > prayer > sermon).
    """
    cats = item.get("categories", {})
    for cat in _cat_order():
        cat_data = cats.get(cat)
        if not cat_data or not cat_data.get("found"):
            continue
        ranges = cat_data.get("ranges") or []
        if any(s <= slide_idx <= e for s, e in ranges):
            return cat
    return None


def _count_category_slides(item: dict) -> Dict[str, int]:
    """Возвращает {cat: slide_count, "other": n, "total": total}."""
    total = item.get("total_slides", 0)
    counts = {cat: 0 for cat in _cat_order()}
    counts["other"] = 0
    counts["total"] = total

    if total <= 0:
        return counts

    for slide_idx in range(1, total + 1):
        cat = _classify_slide(item, slide_idx)
        if cat is None:
            counts["other"] += 1
        else:
            counts[cat] += 1

    return counts


def _get_default_selected_categories(
    item: dict,
    user_config: dict,
) -> set:
    """Возвращает set категорий, отмеченных по умолчанию."""
    selected = set()
    cats = item.get("categories", {})
    meta = _cat_meta()

    for cat in _cat_order():
        cat_data = cats.get(cat)
        if not cat_data or not cat_data.get("found"):
            continue
        flag = meta.get(cat, {}).get("settings_flag", f"process_{cat}")
        if user_config.get(flag, True):
            selected.add(cat)

    counts = _count_category_slides(item)
    if user_config.get("process_other", False) and counts.get("other", 0) > 0:
        selected.add("other")

    return selected


# ==========================================
# СПИННЕР ПРОГРЕССА
# ==========================================

_MODE_LABELS = {
    "sermon": "🎯 Только проповедь",
    "other":  "📄 Только остальные",
    "both":   "📦 Проповедь + остальные",
}

_active_spinners: dict[str, tuple[asyncio.Event, asyncio.Task]] = {}


def _mode_label(mode: str) -> str:
    if mode == "skip":
        return "⏭ Пропущено"
    meta = _cat_meta().get(mode)
    if meta:
        return f"{meta['emoji']} {meta['name']}"
    return _MODE_LABELS.get(mode, f"❓ {mode}")


async def _yd_progress_spinner(
    status_msg,
    task_id: str,
    base_text: str,
    stop_event: asyncio.Event,
    interval: float = 1.2,
) -> None:
    """Циклически дописывает '.', '..', '...' в конец статусного сообщения."""
    dots = [".", "..", "..."]
    idx = 0
    while not stop_event.is_set():
        try:
            await asyncio.wait_for(stop_event.wait(), timeout=interval)
            return
        except asyncio.TimeoutError:
            pass

        if stop_event.is_set():
            return

        try:
            await status_msg.edit_text(
                f"{base_text} {dots[idx]}",
                parse_mode="HTML",
                reply_markup=_get_cancel_keyboard(task_id).as_markup(),
            )
        except asyncio.CancelledError:
            raise
        except TelegramBadRequest as e:
            if "message is not modified" not in str(e):
                logging.debug(f"[YD-SPINNER] TelegramBadRequest: {e}")
        except Exception as e:
            logging.debug(f"[YD-SPINNER] edit_text: {e}")

        idx = (idx + 1) % len(dots)


async def _yd_stop_spinner(task_id: str, wait_timeout: float = 2.0) -> None:
    """Останавливает активный спиннер задачи."""
    entry = _active_spinners.pop(task_id, None)
    if entry is None:
        return
    stop_event, spinner_task = entry
    stop_event.set()
    try:
        await asyncio.wait_for(asyncio.shield(spinner_task), timeout=wait_timeout)
    except (asyncio.TimeoutError, asyncio.CancelledError):
        spinner_task.cancel()
        try:
            await spinner_task
        except (asyncio.CancelledError, Exception):
            pass
    except Exception as e:
        logging.debug(f"[YD-SPINNER] stop: {e}")


async def _yd_with_spinner(status_msg, task_id: str, base_text: str, coro):
    """Запускает корутину, параллельно анимируя статусное сообщение."""
    stop_event = asyncio.Event()
    spinner_task = asyncio.create_task(
        _yd_progress_spinner(status_msg, task_id, base_text, stop_event)
    )
    _active_spinners[task_id] = (stop_event, spinner_task)
    try:
        return await coro
    finally:
        entry = _active_spinners.get(task_id)
        if entry is not None and entry[1] is spinner_task:
            _active_spinners.pop(task_id, None)
        stop_event.set()
        try:
            await asyncio.wait_for(asyncio.shield(spinner_task), timeout=2.0)
        except (asyncio.TimeoutError, asyncio.CancelledError):
            spinner_task.cancel()
            try:
                await spinner_task
            except (asyncio.CancelledError, Exception):
                pass
        except Exception as e:
            logging.debug(f"[YD-SPINNER] finalize: {e}")


# ==========================================
# ЗАЩИТА IN-FLIGHT ОПЕРАЦИЙ
# ==========================================

_active_ops: dict[str, set[asyncio.Task]] = {}
_deferred_cleanup_tasks: set[asyncio.Task] = set()
_deferred_dirs: set[Path] = set()
_cleaning_tasks: set[str] = set()
_detached_workers: set[asyncio.Task] = set()


async def _yd_run_protected(task_id: str, coro):
    """Запускает корутину как Task, регистрирует в _active_ops и защищает shield."""
    async def _runner():
        try:
            return await coro
        finally:
            bucket = _active_ops.get(task_id)
            if bucket is not None:
                bucket.discard(asyncio.current_task())
                if not bucket:
                    _active_ops.pop(task_id, None)

    op_task = asyncio.create_task(_runner())
    _active_ops.setdefault(task_id, set()).add(op_task)
    return await asyncio.shield(op_task)


async def _yd_to_thread(task_id: str, fn, *args, **kwargs):
    """asyncio.to_thread + регистрация в _active_ops."""
    return await _yd_run_protected(
        task_id, asyncio.to_thread(fn, *args, **kwargs)
    )


async def _yd_async_protected(task_id: str, coro):
    """Защита сетевых корутин Яндекса от внешней отмены."""
    async def _runner():
        try:
            return await coro
        finally:
            bucket = _active_ops.get(task_id)
            if bucket is not None:
                bucket.discard(asyncio.current_task())
                if not bucket:
                    _active_ops.pop(task_id, None)

    op_task = asyncio.create_task(_runner())
    _active_ops.setdefault(task_id, set()).add(op_task)
    return await asyncio.shield(op_task)


async def _yd_deferred_task_dir_cleanup(task_dir: Path, ops: list) -> None:
    """Фоновое удаление task_dir: ждём завершения in-flight ops."""
    try:
        await asyncio.gather(*ops, return_exceptions=True)
    except Exception as e:
        logging.debug(f"[YD-DEFERRED] gather: {e}")
    _safe_delete_task_dir(task_dir)


def _spawn_deferred_cleanup(task_dir: Path, ops: list) -> None:
    """Создаёт deferred cleanup task с защитой от дублирования."""
    if task_dir is None:
        return
    if task_dir in _deferred_dirs:
        logging.info(f"[YD-DEFERRED] {task_dir} уже в очереди — skip")
        return
    _deferred_dirs.add(task_dir)

    task = asyncio.create_task(_yd_deferred_task_dir_cleanup(task_dir, ops))
    _deferred_cleanup_tasks.add(task)

    def _done(t: asyncio.Task):
        _deferred_cleanup_tasks.discard(t)
        _deferred_dirs.discard(task_dir)

    task.add_done_callback(_done)


async def _cancel_worker_async(
    worker: asyncio.Task,
    task_id: str,
    timeout: float = 15.0,
) -> None:
    """Фоновая отмена worker'а."""
    if worker is None or worker.done() or worker is asyncio.current_task():
        return
    try:
        worker.cancel()
    except Exception as e:
        logging.debug(f"[YD-CANCEL-ASYNC] cancel {task_id}: {e}")
    try:
        await asyncio.wait_for(asyncio.shield(worker), timeout=timeout)
        logging.info(f"[YD-CANCEL-ASYNC] worker {task_id} завершён")
    except asyncio.TimeoutError:
        logging.warning(
            f"[YD-CANCEL-ASYNC] worker {task_id} не завершился за {timeout}s"
        )
    except asyncio.CancelledError:
        pass
    except Exception as e:
        logging.debug(f"[YD-CANCEL-ASYNC] await {task_id}: {e}")


def _spawn_cancel_worker(worker: asyncio.Task, task_id: str) -> None:
    """Запускает фоновую отмену worker'а."""
    t = asyncio.create_task(_cancel_worker_async(worker, task_id))
    _deferred_cleanup_tasks.add(t)
    t.add_done_callback(_deferred_cleanup_tasks.discard)


async def _run_worker_detached(
    task_id: str,
    coro_factory: Callable[[], Awaitable[None]],
    *,
    name: str = "yd_worker",
) -> asyncio.Task:
    """Запускает worker в отдельном task'е, регистрирует в session, возвращает сразу."""
    async def _wrapper():
        try:
            await coro_factory()
        except asyncio.CancelledError:
            logging.info(f"[YD-WORKER] {name} {task_id} отменён")
        except Exception as e:
            logging.error(
                f"[YD-WORKER] {name} {task_id} упал: {e}", exc_info=True
            )

    worker = asyncio.create_task(_wrapper())
    _detached_workers.add(worker)

    def _done(t: asyncio.Task):
        _detached_workers.discard(t)
        sess = sessions.get(task_id)
        if sess is not None and sess.get("worker_task") is t:
            sess["worker_task"] = None

    worker.add_done_callback(_done)

    async with yd_session_lock:
        sess = sessions.get(task_id)
        if sess is not None:
            sess["worker_task"] = worker

    return worker


async def drain_deferred_cleanups(timeout: float = 30.0) -> None:
    """Дожидается завершения отложенных cleanup-задач и detached worker'ов."""
    pending_cleanups = [t for t in _deferred_cleanup_tasks if not t.done()]
    pending_workers = [t for t in _detached_workers if not t.done()]

    if not pending_cleanups and not pending_workers:
        return

    logging.info(
        f"[YD-DRAIN] Ожидаю {len(pending_cleanups)} cleanup-задач "
        f"и {len(pending_workers)} detached worker'ов..."
    )
    try:
        await asyncio.wait_for(
            asyncio.gather(
                *pending_cleanups, *pending_workers,
                return_exceptions=True,
            ),
            timeout=timeout,
        )
        logging.info("[YD-DRAIN] Все отложенные задачи завершены")
    except asyncio.TimeoutError:
        logging.warning(
            f"[YD-DRAIN] Не все задачи завершились за {timeout}s"
        )


# ==========================================
# ПАЙПЛАЙН ПОДГОТОВКИ
# ==========================================

async def _yd_prepare_files_impl(
    callback: types.CallbackQuery,
    bot: Bot,
    SHM_DIR: str,
    user_mgr,
    files_to_process: list,
    sunday,
    paths: dict,
    session_key: str,
    nonce: str,
):
    """Первая фаза: скачивание + извлечение заметок + детект категорий."""
    status_msg = callback.message
    owner_user_id = callback.from_user.id
    chat_id = callback.message.chat.id

    task_id = f"yd_task_{secrets.token_hex(6)}"
    task_dir = Path(SHM_DIR) / task_id

    registration_status = "ok"

    async with yd_session_lock:
        picker = sessions.get(session_key)
        if picker is None:
            registration_status = "no_picker"
        elif picker.get("nonce") != nonce:
            logging.warning(
                f"[YD-PREP] picker.nonce={picker.get('nonce')!r} ≠ "
                f"expected nonce={nonce!r} — прерываем"
            )
            registration_status = "wrong_nonce"
        else:
            picker.setdefault("task_ids", []).append(task_id)
            picker["processing"] = False
            sessions[task_id] = {
                "user_id": owner_user_id,
                "chat_id": chat_id,
                "session_key": session_key,
                "pending": None,
                "nonce": nonce,
                "created_at": time.time(),
                "cancelled": False,
                "worker_task": None,
            }
            yd_active_tasks.add(task_id)

    if registration_status == "no_picker":
        await _safe_edit(status_msg, "❌ Сессия была отменена.")
        return

    if registration_status == "wrong_nonce":
        await _safe_edit(
            status_msg,
            "❌ Сессия была заменена другой командой.\n"
            "Запустите /sunday заново."
        )
        return

    try:
        task_dir.mkdir(parents=True, exist_ok=True)
    except Exception as e:
        logging.error(
            f"[YD-PREP] mkdir упал для {task_id}: {e}", exc_info=True,
        )
        try:
            async with yd_session_lock:
                picker = sessions.get(session_key)
                if picker is not None:
                    task_ids = picker.get("task_ids")
                    if isinstance(task_ids, list) and task_id in task_ids:
                        task_ids.remove(task_id)
                sessions.pop(task_id, None)
                yd_active_tasks.discard(task_id)
        except Exception as cleanup_err:
            logging.error(
                f"[YD-PREP] не удалось откатить регистрацию: {cleanup_err}",
                exc_info=True,
            )
        _safe_delete_task_dir(task_dir)
        await _safe_edit(
            status_msg,
            f"❌ Не удалось создать папку задачи: "
            f"<code>{html_module.escape(str(e)[:200])}</code>",
        )
        return

    status_message_id = status_msg.message_id if status_msg is not None else None

    logging.info(
        f"[YD-PREP] Старт: task_id={task_id}, файлов={len(files_to_process)}"
    )

    cleanup_done = False
    try:
        target_base = paths["target"]
        user_config = user_mgr.get_user_config(owner_user_id)
        quality = user_config.get("quality", "2k")

        prepared = []

        for f_idx, pptx_item in enumerate(files_to_process, start=1):
            _touch_task(task_dir)

            task_sess = sessions.get(task_id)
            if task_sess is None or task_sess.get("cancelled"):
                logging.info(
                    f"[YD-PREP] Задача {task_id} отменена — прерываем"
                )
                cleanup_done = True
                await asyncio.shield(_yd_cleanup_task(
                    task_id, session_key, task_dir,
                    owner_user_id, chat_id, nonce,
                ))
                return

            file_name = pptx_item["name"]
            file_name_esc = html_module.escape(file_name)
            file_size = pptx_item.get("size", 0)

            logging.info(
                f"[YD-PREP] Файл #{f_idx}/{len(files_to_process)}: "
                f"{file_name!r} ({_format_size(file_size)})"
            )

            per_file_dir = task_dir / f"src_{f_idx}"
            per_file_dir.mkdir(exist_ok=True)

            base_dl_text = f"📥 Скачиваю <code>{file_name_esc}</code>"
            await _safe_edit(
                status_msg,
                base_dl_text,
                parse_mode="HTML",
                reply_markup=_get_cancel_keyboard(task_id).as_markup(),
            )

            local_pptx = per_file_dir / file_name
            ok = await _yd_with_spinner(
                status_msg,
                task_id,
                base_dl_text,
                _yd_async_protected(
                    task_id,
                    yandex_state.config.client.download_file(
                        pptx_item["path"], local_pptx
                    ),
                ),
            )
            if not ok:
                logging.error(f"[YD-PREP] {file_name}: скачивание провалилось")
                prepared.append({
                    "file_name": file_name,
                    "file_slug": safe_folder_name(file_name),
                    "failed_at_stage": "download",
                })
                continue

            logging.debug(
                f"[YD-PREP] {file_name}: скачан, "
                f"размер={_format_size(local_pptx.stat().st_size)}"
            )

            task_sess = sessions.get(task_id)
            if task_sess is None or task_sess.get("cancelled"):
                logging.info(
                    f"[YD-PREP] Задача {task_id} отменена после скачивания — прерываем"
                )
                cleanup_done = True
                await asyncio.shield(_yd_cleanup_task(
                    task_id, session_key, task_dir,
                    owner_user_id, chat_id, nonce,
                ))
                return

            if local_pptx.suffix.lower() == ".ppt":
                base_norm_text = (
                    f"🔧 Готовлю .ppt → .pptx: <code>{file_name_esc}</code>"
                )
                await _safe_edit(
                    status_msg,
                    base_norm_text,
                    parse_mode="HTML",
                    reply_markup=_get_cancel_keyboard(task_id).as_markup(),
                )
                try:
                    normalized_pptx = await _yd_with_spinner(
                        status_msg,
                        task_id,
                        base_norm_text,
                        _yd_to_thread(
                            task_id,
                            ppt_to_pptx_crossplatform,
                            local_pptx, per_file_dir,
                        ),
                    )
                except Exception as e:
                    logging.error(
                        f"[YD-PREP] {file_name}: .ppt→.pptx упал: {e}",
                        exc_info=True,
                    )
                    prepared.append({
                        "file_name": file_name,
                        "file_slug": safe_folder_name(file_name),
                        "failed_at_stage": "normalize",
                    })
                    continue

                if not normalized_pptx or not Path(normalized_pptx).exists():
                    logging.error(f"[YD-PREP] {file_name}: .pptx не создан")
                    prepared.append({
                        "file_name": file_name,
                        "file_slug": safe_folder_name(file_name),
                        "failed_at_stage": "normalize",
                    })
                    continue
            else:
                normalized_pptx = local_pptx

            task_sess = sessions.get(task_id)
            if task_sess is None or task_sess.get("cancelled"):
                logging.info(
                    f"[YD-PREP] Задача {task_id} отменена после нормализации — прерываем"
                )
                cleanup_done = True
                await asyncio.shield(_yd_cleanup_task(
                    task_id, session_key, task_dir,
                    owner_user_id, chat_id, nonce,
                ))
                return

            total_slides = None
            try:
                from pptx import Presentation
                prs = Presentation(str(normalized_pptx))
                total_slides = len(prs.slides._sldIdLst)
            except Exception as e:
                logging.warning(
                    f"[YD-PREP] {file_name}: python-pptx не открыл: {e}"
                )

            if not total_slides:
                renorm_result = await _yd_to_thread(
                    task_id,
                    librenormalize_to_pptx, normalized_pptx, per_file_dir,
                )
                if renorm_result is None:
                    logging.error(
                        f"[YD-PREP] {file_name}: файл не читается"
                    )
                    prepared.append({
                        "file_name": file_name,
                        "file_slug": safe_folder_name(file_name),
                        "failed_at_stage": "count",
                    })
                    continue
                normalized_pptx, total_slides = renorm_result
                logging.info(
                    f"[YD-PREP] {file_name}: нормализован через LibreOffice "
                    f"({total_slides} слайдов)"
                )

            if not total_slides:
                logging.error(
                    f"[YD-PREP] {file_name}: total_slides == 0, пропускаем"
                )
                prepared.append({
                    "file_name": file_name,
                    "file_slug": safe_folder_name(file_name),
                    "failed_at_stage": "count",
                })
                continue

            await _safe_edit(
                status_msg,
                f"🔍 Читаю заметки докладчика: <code>{file_name_esc}</code>...",
                parse_mode="HTML",
                reply_markup=_get_cancel_keyboard(task_id).as_markup(),
            )

            notes_ok, notes, incomplete = await _yd_to_thread(
                task_id, extract_speaker_notes, str(normalized_pptx)
            )

            categories = _detect_all_categories(
                notes, notes_ok, incomplete
            )

            if not notes_ok:
                incomplete_warning = (
                    "⚠️ Не удалось прочитать заметки докладчика."
                )
            elif incomplete:
                incomplete_warning = (
                    "⚠️ Заметки прочитаны частично, "
                    "категории не определены автоматически."
                )
            else:
                incomplete_warning = None

            logging.info(
                f"[YD-PREP] {file_name}: заметок={len(notes)}, "
                f"categories="
                + ", ".join(
                    f"{c}({categories[c]['found']})"
                    for c in _cat_order()
                )
                + f", notes_ok={notes_ok}, incomplete={incomplete}"
            )

            sermon_data = categories.get("sermon") or {}
            legacy_ranges = sermon_data.get("ranges")
            legacy_start, legacy_end = _ranges_to_start_end(legacy_ranges)

            item = {
                "file_name": file_name,
                "file_slug": safe_folder_name(file_name),
                "file_path": normalized_pptx,
                "total_slides": total_slides,
                "categories": categories,
                "start": legacy_start,
                "end": legacy_end,
                "ranges": legacy_ranges,
                "matches": sermon_data.get("matches", []),
                "notes_ok": notes_ok,
                "incomplete": incomplete,
                "incomplete_warning": incomplete_warning,
                "confirmed": False,
                "convert_mode": None,
                "selected_categories": None,
                "_user_config": dict(user_config),
            }

            item["selected_categories"] = _get_default_selected_categories(
                item, user_config
            )

            prepared.append(item)

        cleanup_needed = False
        async with yd_session_lock:
            picker = sessions.get(session_key)
            if picker is None or picker.get("cancelled"):
                cleanup_needed = True
            else:
                task_sess = sessions.get(task_id)
                if task_sess is None or task_sess.get("cancelled"):
                    cleanup_needed = True
                else:
                    task_sess["pending"] = {
                        "task_dir": task_dir,
                        "target_base": target_base,
                        "quality": quality,
                        "prepared": prepared,
                        "owner_user_id": owner_user_id,
                        "chat_id": chat_id,
                        "session_key": session_key,
                        "nonce": nonce,
                        "bot": bot,
                        "prompt_nonce": None,
                        "prompt_idx": None,
                        "prompt_message_id": None,
                        "prompt_timeout_task": None,
                        "prompt_watchdog_nonce": None,
                        "status_message_id": status_message_id,
                        "awaiting_range_for_idx": None,
                    }

        if cleanup_needed:
            logging.info(
                f"[YD-PREP] Задача {task_id} отменена до сохранения pending"
            )
            cleanup_done = True
            await asyncio.shield(_yd_cleanup_task(
                task_id, session_key, task_dir,
                owner_user_id, chat_id, nonce,
            ))
            return

        needs_confirm = [
            p for p in prepared
            if "file_path" in p and not p.get("confirmed")
        ]

        logging.info(
            f"[YD-PREP] Готово: prepared={len(prepared)}, "
            f"needs_confirm={len(needs_confirm)}"
        )

        if not needs_confirm:
            for item in prepared:
                if not item.get("selected_categories"):
                    item["selected_categories"] = (
                        {"other"} if "file_path" in item else set()
                    )
                item["confirmed"] = True
            await _run_worker_detached(
                task_id,
                lambda: _yd_convert_and_upload_impl(
                    bot=bot,
                    task_id=task_id,
                    status_msg=status_msg,
                ),
                name="convert",
            )
            return

        await _yd_render_category_prompt(
            task_id=task_id,
            item=needs_confirm[0],
            status_msg=status_msg,
        )

    except asyncio.CancelledError:
        logging.info(
            f"[YD-PREP] Worker {task_id} отменён — выполняю cleanup"
        )
        cleanup_done = True
        try:
            await asyncio.shield(_yd_cleanup_task(
                task_id, session_key, task_dir,
                owner_user_id, chat_id, nonce,
            ))
        except Exception as e:
            logging.error(
                f"[YD-PREP] cleanup после отмены упал: {e}", exc_info=True
            )
        raise

    except Exception as e:
        logging.error(f"[YD-PREP] Ошибка: {e}", exc_info=True)
        cleanup_done = True
        await _yd_cleanup_task(
            task_id, session_key, task_dir,
            owner_user_id, chat_id, nonce,
            bot=bot, status_msg=status_msg, error=e,
        )


async def _yd_prepare_files(
    callback: types.CallbackQuery,
    bot: Bot,
    SHM_DIR: str,
    user_mgr,
    files_to_process: list,
    sunday,
    paths: dict,
    session_key: str,
    nonce: str,
) -> asyncio.Task:
    """Запускает _yd_prepare_files_impl как detached worker."""
    temp_task_id = f"yd_pending_{secrets.token_hex(4)}"

    async def _factory():
        await _yd_prepare_files_impl(
            callback=callback,
            bot=bot,
            SHM_DIR=SHM_DIR,
            user_mgr=user_mgr,
            files_to_process=files_to_process,
            sunday=sunday,
            paths=paths,
            session_key=session_key,
            nonce=nonce,
        )

    return await _run_worker_detached(
        temp_task_id, _factory, name="prepare",
    )


# ==========================================
# ПРОМПТЫ
# ==========================================

async def _yd_claim_prompt(callback: types.CallbackQuery) -> Optional[dict]:
    """Атомарно проверяет и 'потребляет' промпт."""
    parts = callback.data.split(":")
    if len(parts) < 4:
        await _safe_answer(callback, "❌ Некорректный запрос.", show_alert=True)
        return None
    task_id = parts[1]
    try:
        idx = int(parts[2])
    except ValueError:
        await _safe_answer(callback, "❌ Некорректный запрос.", show_alert=True)
        return None
    nonce = parts[3]

    session = sessions.get(task_id)
    if not session or "pending" not in session:
        await _safe_answer(callback, "❌ Сессия неактивна.", show_alert=True)
        return None
    if session.get("cancelled"):
        await _safe_answer(callback, "❌ Задача отменена.", show_alert=True)
        return None

    pending = session["pending"]
    if not isinstance(pending, dict):
        await _safe_answer(callback, "❌ Сессия неактивна.", show_alert=True)
        return None

    if pending.get("prompt_nonce") != nonce or pending.get("prompt_idx") != idx:
        await _safe_answer(callback, "⏳ Промпт уже обработан.", show_alert=True)
        return None

    if callback.from_user.id != pending["owner_user_id"]:
        await _safe_answer(callback, "❌ Только автор.", show_alert=True)
        return None

    timeout_task = pending.get("prompt_timeout_task")
    if timeout_task is not None and not timeout_task.done():
        timeout_task.cancel()
    pending["prompt_timeout_task"] = None
    pending["prompt_watchdog_nonce"] = None

    pending["prompt_nonce"] = None
    pending["prompt_idx"] = None
    pending["prompt_message_id"] = None
    pending["awaiting_range_for_idx"] = None

    try:
        await callback.message.edit_reply_markup(reply_markup=None)
    except TelegramBadRequest as e:
        if "message is not modified" not in str(e):
            logging.debug(f"claim_prompt edit_reply_markup: {e}")
    except Exception as e:
        logging.debug(f"claim_prompt edit_reply_markup: {e}")

    return pending


def _render_category_toggle_keyboard(
    task_id: str, idx: int, nonce: str,
    item: dict,
    selected: set,
    total_slides: int,
) -> InlineKeyboardBuilder:
    """Клавиатура с toggle-кнопками категорий."""
    kb = InlineKeyboardBuilder()
    counts = _count_category_slides(item)
    cats = item.get("categories", {})
    meta = _cat_meta()

    for cat in _cat_order():
        cat_meta = meta.get(cat, {})
        emoji = cat_meta.get("emoji", "❓")
        name = cat_meta.get("name", cat)
        cat_data = cats.get(cat) or {}
        found = bool(cat_data.get("found"))
        count = counts.get(cat, 0)

        if not found:
            kb.row(InlineKeyboardButton(
                text=f"⚫ {emoji} {name} — не найдено",
                callback_data=f"yd_cat_noop:{task_id}:{idx}:{nonce}:{cat}",
            ))
            continue

        ranges = cat_data.get("ranges") or []
        ranges_text = _format_ranges_text(ranges)
        check = "✅" if cat in selected else "⬜"
        if count > 0:
            text = f"{check} {emoji} {name} ({ranges_text}, {count} с.)"
        else:
            text = f"{check} {emoji} {name} ({ranges_text})"

        kb.row(InlineKeyboardButton(
            text=text,
            callback_data=f"yd_cat_toggle:{task_id}:{idx}:{nonce}:{cat}",
        ))

    other_count = counts.get("other", 0)
    if other_count > 0:
        check = "✅" if "other" in selected else "⬜"
        kb.row(InlineKeyboardButton(
            text=f"{check} 📄 Остальные ({other_count} с.)",
            callback_data=f"yd_cat_toggle:{task_id}:{idx}:{nonce}:other",
        ))

    kb.row(InlineKeyboardButton(
        text="✏️ Изменить диапазон категории",
        callback_data=f"yd_cat_edit:{task_id}:{idx}:{nonce}",
    ))

    total_selected = 0
    for cat in _cat_order():
        if cat in selected:
            total_selected += counts.get(cat, 0)
    if "other" in selected:
        total_selected += other_count

    if total_selected > 0:
        kb.row(InlineKeyboardButton(
            text=f"✅ Конвертировать выбранные ({total_selected} с.)",
            callback_data=f"yd_cat_convert:{task_id}:{idx}:{nonce}",
        ))
    else:
        kb.row(InlineKeyboardButton(
            text="✅ Конвертировать выбранные",
            callback_data=f"yd_cat_noop:{task_id}:{idx}:{nonce}:empty",
        ))

    kb.row(InlineKeyboardButton(
        text="❌ Отменить задачу",
        callback_data=f"yd_task_cancel:{task_id}",
    ))

    return kb


def _render_category_prompt_text(item: dict, selected: set) -> str:
    """Формирует текст промпта категорий."""
    file_name = item.get("file_name", "")
    file_esc = html_module.escape(file_name)
    total_slides = item.get("total_slides", 0)
    cats = item.get("categories", {})
    meta = _cat_meta()
    counts = _count_category_slides(item)

    lines = [
        f"🎬 <b>Выбор категорий для конвертации</b>",
        "",
        f"📄 Файл: <code>{file_esc}</code>",
        f"📌 Всего слайдов: <b>{total_slides}</b>",
        "",
    ]

    found_lines = []
    for cat in _cat_order():
        cat_data = cats.get(cat) or {}
        if not cat_data.get("found"):
            continue
        cat_meta = meta.get(cat, {})
        emoji = cat_meta.get("emoji", "❓")
        name = cat_meta.get("name", cat)
        count = counts.get(cat, 0)
        ranges = cat_data.get("ranges") or []
        ranges_text = _format_ranges_text(ranges)
        manual = " (вручную)" if cat_data.get("manual") else ""
        found_lines.append(
            f"{emoji} {name}{manual}: <b>{ranges_text}</b> ({count} с.)"
        )

    if found_lines:
        lines.extend(found_lines)
        lines.append("")
    else:
        lines.append("⚠️ <i>Ни одна категория не найдена автоматически.</i>")
        lines.append("")

    other_count = counts.get("other", 0)
    if other_count > 0:
        lines.append(f"📄 Остальные: <b>{other_count} с.</b>")
        lines.append("")

    lines.append("🎯 <b>Отметьте категории для конвертации:</b>")

    if not selected:
        lines.append("")
        lines.append(
            "⚠️ <i>Ничего не выбрано. Отметьте хотя бы одну категорию.</i>"
        )

    return "\n".join(lines)


async def _yd_render_category_prompt(
    task_id: str,
    item: dict,
    status_msg,
    reply_fn=None,
) -> None:
    """Рендер нового промпта категорий."""
    session = sessions.get(task_id)
    if not session or "pending" not in session:
        return
    if session.get("cancelled"):
        return
    pending = session["pending"]
    if not isinstance(pending, dict):
        return

    try:
        idx = pending["prepared"].index(item)
    except (ValueError, KeyError):
        logging.error(
            f"_yd_render_category_prompt: item не найден для {task_id}"
        )
        return

    pending["prompt_idx"] = idx

    if pending.get("prompt_nonce") is None:
        pending["prompt_nonce"] = secrets.token_hex(4)
    prompt_nonce = pending["prompt_nonce"]

    total_slides = item.get("total_slides", 0)
    if total_slides == 0:
        file_path = item.get("file_path")
        if file_path and Path(file_path).exists():
            try:
                from pptx import Presentation
                prs = Presentation(str(file_path))
                total_slides = len(prs.slides._sldIdLst)
                item["total_slides"] = total_slides
            except Exception as e:
                logging.warning(
                    f"Не удалось определить число слайдов: {e}"
                )

    if item.get("selected_categories") is None:
        item["selected_categories"] = _get_default_selected_categories(
            item, item.get("_user_config", {})
        )
    selected = set(item.get("selected_categories") or set())

    text = _render_category_prompt_text(item, selected)
    kb = _render_category_toggle_keyboard(
        task_id, idx, prompt_nonce, item, selected, total_slides
    )

    sent_msg = None
    if reply_fn is not None:
        sent_msg = await reply_fn(
            text, parse_mode="HTML", reply_markup=kb.as_markup()
        )
    elif status_msg is not None:
        await status_msg.edit_text(
            text, parse_mode="HTML", reply_markup=kb.as_markup()
        )
        sent_msg = status_msg
    else:
        logging.warning(
            f"_yd_render_category_prompt: нет ни reply_fn, ни status_msg "
            f"для {task_id}"
        )
        return

    if pending.get("prompt_nonce") != prompt_nonce:
        logging.info(
            f"_yd_render_category_prompt: nonce изменён конкурентно"
        )
        return

    if sent_msg is not None and hasattr(sent_msg, "message_id"):
        pending["prompt_message_id"] = sent_msg.message_id

    existing_task = pending.get("prompt_timeout_task")
    existing_nonce = pending.get("prompt_watchdog_nonce")

    if existing_task is not None and not existing_task.done():
        if existing_nonce == prompt_nonce:
            return
        existing_task.cancel()
        # ✅ v4.0.1: дожидаемся реальной отмены — старый watchdog
        # может успеть дойти до _yd_cleanup_task.
        try:
            await asyncio.wait_for(
                asyncio.shield(existing_task), timeout=5.0
            )
        except (asyncio.TimeoutError, asyncio.CancelledError):
            pass
        except Exception as e:
            logging.debug(f"render_category_prompt: await old: {e}")
        pending["prompt_timeout_task"] = None
        pending["prompt_watchdog_nonce"] = None

    pending["prompt_timeout_task"] = asyncio.create_task(
        _yd_prompt_timeout_watchdog(
            task_id, yandex_state.config.prompt_timeout_sec, prompt_nonce
        )
    )
    pending["prompt_watchdog_nonce"] = prompt_nonce


async def _yd_render_sermon_prompt(
    task_id: str,
    item: dict,
    status_msg,
    reply_fn=None,
) -> None:
    """Старый промпт (fallback для yd_sermon_edit и handlers.py)."""
    session = sessions.get(task_id)
    if not session or "pending" not in session:
        return
    if session.get("cancelled"):
        return
    pending = session["pending"]
    if not isinstance(pending, dict):
        return

    try:
        idx = pending["prepared"].index(item)
    except (ValueError, KeyError):
        logging.error(
            f"_yd_render_sermon_prompt: item не найден для {task_id}"
        )
        return

    pending["prompt_idx"] = idx

    if pending.get("prompt_nonce") is None:
        pending["prompt_nonce"] = secrets.token_hex(4)
    prompt_nonce = pending["prompt_nonce"]

    _normalize_item_ranges(item)

    matches = item.get("matches", []) or []
    start = item.get("start")
    end = item.get("end")
    ranges = item.get("ranges")
    file_name = item.get("file_name", "")
    file_esc = html_module.escape(file_name)

    total_slides = item.get("total_slides", 0)
    if total_slides == 0:
        file_path = item.get("file_path")
        if file_path and Path(file_path).exists():
            try:
                from pptx import Presentation
                prs = Presentation(str(file_path))
                total_slides = len(prs.slides._sldIdLst)
            except Exception as e:
                logging.warning(f"Не удалось определить число слайдов: {e}")

    kb = InlineKeyboardBuilder()

    has_valid_range = (
        bool(ranges)
        or (start is not None and end is not None and start <= end)
    )

    if has_valid_range:
        if ranges:
            sermon_count = _count_slides_in_ranges(ranges, total_slides)
            ranges_text = _format_ranges_text(ranges, start, end)
        else:
            sermon_count = _count_slides_for_range(start, end, total_slides)
            ranges_text = _format_ranges_text(None, start, end)

        other_count = max(0, total_slides - sermon_count)

        manual_range = item.get("manual_range", False)

        if manual_range:
            text = (
                f"🎯 <b>Диапазон проповеди установлен</b>\n\n"
                f"📄 Файл: <code>{file_esc}</code>\n"
                f"📊 Диапазон: <b>{ranges_text}</b>\n"
                f"🎯 Слайдов проповеди: <b>{sermon_count}</b>\n\n"
                f"❓ <b>Какие слайды конвертировать?</b>"
            )
        else:
            preview = ", ".join(str(n) for n in matches[:15])
            if len(matches) > 15:
                preview += f" …и ещё {len(matches) - 15}"

            text = (
                f"🎯 <b>Найдена пометка «проповедь»</b>\n\n"
                f"📄 Файл: <code>{file_esc}</code>\n"
                f"📌 Слайды с пометкой: <code>{preview}</code>\n"
                f"📊 Предлагаемый диапазон: <b>{ranges_text}</b>\n\n"
                f"❓ <b>Какие слайды конвертировать?</b>"
            )

        kb.row(
            InlineKeyboardButton(
                text=f"✅ Только проповедь ({sermon_count} с.)",
                callback_data=(
                    f"yd_sermon_mode:{task_id}:{idx}:{prompt_nonce}:sermon"
                ),
            ),
        )
        kb.row(
            InlineKeyboardButton(
                text=f"📄 Только остальные ({other_count} с.)",
                callback_data=(
                    f"yd_sermon_mode:{task_id}:{idx}:{prompt_nonce}:other"
                ),
            ),
        )
        kb.row(
            InlineKeyboardButton(
                text=f"📦 Проповеди и остальное раздельно ({total_slides} с.)",
                callback_data=(
                    f"yd_sermon_mode:{task_id}:{idx}:{prompt_nonce}:both"
                ),
            ),
        )
        kb.row(
            InlineKeyboardButton(
                text="✏️ Изменить диапазон Проповеди",
                callback_data=f"yd_sermon_edit:{task_id}:{idx}:{prompt_nonce}",
            ),
        )
    else:
        notes_ok = item.get("notes_ok", True)
        incomplete = item.get("incomplete", False)

        if not notes_ok:
            reason = (
                "⚠️ <b>Не удалось прочитать заметки докладчика.</b>\n"
                "Возможно, файл повреждён или содержит только изображения."
            )
        elif incomplete:
            reason = (
                "⚠️ <b>Заметки прочитаны частично.</b>\n"
                "Автоматически определить проповедь не удалось."
            )
        elif matches:
            match_str = ", ".join(str(n) for n in matches[:5])
            reason = (
                f"📌 Найдена <b>одна</b> пометка «проповедь»: "
                f"слайд <code>{match_str}</code>\n"
                f"Для одной пометки авто-диапазон не строится."
            )
        else:
            reason = "В заметках докладчика нет слова «проповедь»."

        text = (
            f"🤔 <b>Автоматически определить проповедь не удалось</b>\n\n"
            f"📄 Файл: <code>{file_esc}</code>\n"
            f"📌 Всего слайдов: <b>{total_slides}</b>\n\n"
            f"{reason}\n\n"
            f"❓ <b>Что делать с файлом?</b>"
        )

        kb.row(
            InlineKeyboardButton(
                text=f"📄 Конвертировать всё ({total_slides} с.)",
                callback_data=(
                    f"yd_sermon_mode:{task_id}:{idx}:{prompt_nonce}:other"
                ),
            ),
        )
        kb.row(
            InlineKeyboardButton(
                text="✏️ Указать диапазон Проповеди",
                callback_data=f"yd_sermon_edit:{task_id}:{idx}:{prompt_nonce}",
            ),
        )

    kb.row(
        InlineKeyboardButton(
            text="❌ Отменить задачу",
            callback_data=f"yd_task_cancel:{task_id}",
        ),
    )

    sent_msg = None
    if reply_fn is not None:
        sent_msg = await reply_fn(
            text, parse_mode="HTML", reply_markup=kb.as_markup()
        )
    elif status_msg is not None:
        await status_msg.edit_text(
            text, parse_mode="HTML", reply_markup=kb.as_markup()
        )
        sent_msg = status_msg
    else:
        logging.warning(
            f"_yd_render_sermon_prompt: нет ни reply_fn, ни status_msg "
            f"для {task_id}"
        )
        return

    if pending.get("prompt_nonce") != prompt_nonce:
        logging.info(
            f"_yd_render_sermon_prompt: nonce изменён конкурентно"
        )
        return

    if sent_msg is not None and hasattr(sent_msg, "message_id"):
        pending["prompt_message_id"] = sent_msg.message_id

    existing_task = pending.get("prompt_timeout_task")
    existing_nonce = pending.get("prompt_watchdog_nonce")

    if existing_task is not None and not existing_task.done():
        if existing_nonce == prompt_nonce:
            return
        existing_task.cancel()
        pending["prompt_timeout_task"] = None
        pending["prompt_watchdog_nonce"] = None

    pending["prompt_timeout_task"] = asyncio.create_task(
        _yd_prompt_timeout_watchdog(
            task_id, yandex_state.config.prompt_timeout_sec, prompt_nonce
        )
    )
    pending["prompt_watchdog_nonce"] = prompt_nonce


async def _yd_prompt_timeout_watchdog(
    task_id: str, timeout_sec: int, expected_nonce: str
):
    """Если пользователь не ответил — уведомляем и очищаем."""
    try:
        await asyncio.sleep(timeout_sec)

        snapshot = None
        async with yd_session_lock:
            session = sessions.get(task_id)
            if not session or "pending" not in session:
                return
            pending = session["pending"]
            if not isinstance(pending, dict):
                return
            if pending.get("prompt_nonce") != expected_nonce:
                logging.info(
                    f"[YD-PROMPT] watchdog {task_id}: nonce устарел — skip"
                )
                return
            snapshot = {
                "chat_id": pending.get("chat_id"),
                "bot": pending.get("bot"),
                "owner_user_id": pending.get("owner_user_id"),
                "session_key": pending.get("session_key"),
                "task_dir": pending.get("task_dir"),
                "nonce": pending.get("nonce"),
                "prompt_message_id": pending.get("prompt_message_id"),
            }

        if snapshot is None:
            return

        logging.info(
            f"[YD-PROMPT] ⏰ Промпт {task_id} не подтверждён за {timeout_sec}s"
        )

        chat_id = snapshot["chat_id"]
        bot: Optional[Bot] = snapshot["bot"]
        owner_user_id = snapshot["owner_user_id"]
        session_key = snapshot["session_key"]
        task_dir = snapshot["task_dir"]
        nonce = snapshot["nonce"]
        prompt_message_id = snapshot["prompt_message_id"]

        if bot is not None and chat_id is not None:
            minutes = max(1, timeout_sec // 60)
            try:
                await bot.send_message(
                    chat_id=chat_id,
                    text=(
                        f"⏰ <b>Время ожидания истекло</b> ({minutes} мин).\n\n"
                        f"Задача отменена, временные файлы удалены.\n"
                        f"Если хотите обработать файл — запустите /sunday заново."
                    ),
                    parse_mode="HTML",
                )
            except Exception as e:
                logging.warning(f"Не удалось уведомить о таймауте: {e}")

        if bot is not None and chat_id is not None and prompt_message_id is not None:
            try:
                await bot.edit_message_reply_markup(
                    chat_id=chat_id,
                    message_id=prompt_message_id,
                    reply_markup=None,
                )
            except TelegramBadRequest:
                pass
            except Exception:
                pass

        await asyncio.shield(_yd_cleanup_task(
            task_id, session_key, task_dir,
            owner_user_id, chat_id, nonce,
        ))
    except asyncio.CancelledError:
        pass
    except Exception as e:
        logging.error(f"Ошибка watchdog {task_id}: {e}", exc_info=True)


# ==========================================
# ОЧИСТКА
# ==========================================

async def _yd_cleanup_task(
    task_id: str,
    session_key: str,
    task_dir: Path,
    owner_user_id: int,
    chat_id: int,
    nonce: str,
    bot: Optional[Bot] = None,
    status_msg=None,
    error: Optional[Exception] = None,
):
    """Идемпотентная очистка."""
    async with yd_session_lock:
        if task_id in _cleaning_tasks:
            logging.info(
                f"[YD-CLEANUP] task_id={task_id} уже чистится — skip"
            )
            return
        _cleaning_tasks.add(task_id)

    try:
        logging.info(
            f"[YD-CLEANUP] task_id={task_id}, session_key={session_key!r}, "
            f"reason={'error' if error else 'normal'}"
        )

        try:
            await _yd_stop_spinner(task_id)
        except Exception as e:
            logging.debug(f"[YD-CLEANUP] _yd_stop_spinner: {e}")

        try:
            session = sessions.get(task_id)
            if session is not None:
                pending = session.get("pending")
                if isinstance(pending, dict):
                    task_bot = pending.get("bot")
                    task_chat_id = pending.get("chat_id")
                    if task_bot is not None and task_chat_id is not None:
                        message_ids = []
                        for key in ("prompt_message_id", "status_message_id"):
                            mid = pending.get(key)
                            if mid is not None and mid not in message_ids:
                                message_ids.append(mid)

                        for mid in message_ids:
                            try:
                                await task_bot.edit_message_reply_markup(
                                    chat_id=task_chat_id,
                                    message_id=mid,
                                    reply_markup=None,
                                )
                            except TelegramBadRequest as e:
                                if "message is not modified" in str(e):
                                    logging.debug(
                                        f"[YD-CLEANUP] {mid}: клавиатура уже снята"
                                    )
                                else:
                                    logging.debug(
                                        f"[YD-CLEANUP] TelegramBadRequest "
                                        f"с {mid}: {e}"
                                    )
                            except Exception as e:
                                logging.debug(
                                    f"[YD-CLEANUP] Не удалось снять "
                                    f"клавиатуру с {mid}: {e}"
                                )
        except Exception as e:
            logging.debug(f"[YD-CLEANUP] Ошибка снятия клавиатуры: {e}")

        in_flight_ops = _active_ops.get(task_id)
        if in_flight_ops:
            ops_snapshot = list(in_flight_ops)
            logging.info(
                f"[YD-CLEANUP] {len(ops_snapshot)} in-flight ops для {task_id} — "
                f"откладываем удаление task_dir"
            )
            _spawn_deferred_cleanup(task_dir, ops_snapshot)
        else:
            _safe_delete_task_dir(task_dir)

        try:
            session = sessions.get(task_id)
            if session is not None:
                pending = session.get("pending")
                if isinstance(pending, dict):
                    timeout_task = pending.get("prompt_timeout_task")
                    current_task = asyncio.current_task()
                    if (
                        timeout_task is not None
                        and not timeout_task.done()
                        and timeout_task is not current_task
                    ):
                        timeout_task.cancel()
                    pending["prompt_timeout_task"] = None
                    pending["prompt_watchdog_nonce"] = None
                    pending["awaiting_range_for_idx"] = None
                    pending["prompt_nonce"] = None
                    pending["prompt_idx"] = None
        except Exception as e:
            logging.error(
                f"Ошибка отмены watchdog для {task_id}: {e}", exc_info=True
            )

        try:
            async with yd_session_lock:
                picker = sessions.get(session_key)
                if picker is not None and picker.get("nonce") == nonce:
                    picker["processing"] = False
                    task_ids = picker.get("task_ids")
                    if isinstance(task_ids, list) and task_id in task_ids:
                        task_ids.remove(task_id)
                    sessions.pop(session_key, None)
                elif picker is not None:
                    task_ids = picker.get("task_ids")
                    if isinstance(task_ids, list) and task_id in task_ids:
                        task_ids.remove(task_id)
                        logging.warning(
                            f"[YD-CLEANUP] task_id={task_id} оказался в чужом "
                            f"picker'е — удаляем"
                        )

                sessions.pop(task_id, None)
                yd_active_tasks.discard(task_id)
        except Exception as e:
            logging.error(
                f"Ошибка очистки сессий для {task_id}: {e}", exc_info=True
            )

        _active_ops.pop(task_id, None)

        try:
            await asyncio.shield(yd_release(owner_user_id, chat_id, nonce))
        except Exception as e:
            logging.error(f"Ошибка yd_release для {task_id}: {e}", exc_info=True)

        if error is not None and bot is not None and status_msg is not None:
            try:
                await status_msg.edit_text(
                    f"❌ <b>Ошибка обработки</b>\n\n"
                    f"<code>{html_module.escape(str(error)[:200])}</code>\n\n"
                    f"Временные файлы удалены. Попробуйте снова.",
                    parse_mode="HTML",
                )
            except Exception as e:
                logging.error(
                    f"Ошибка отправки сообщения для {task_id}: {e}"
                )
    finally:
        async with yd_session_lock:
            _cleaning_tasks.discard(task_id)


# ==========================================
# КОНВЕРТАЦИЯ + UPLOAD (v4.0 — N+1 ZIP)
# ==========================================

async def _yd_convert_and_upload_impl(
    bot: Bot,
    task_id: str,
    status_msg,
):
    """Вторая фаза v4.0: конвертация PNG + N+1 ZIP + upload."""
    session = sessions.get(task_id)
    if not session or "pending" not in session:
        return

    def _is_cancelled() -> bool:
        s = sessions.get(task_id)
        return s is None or s.get("cancelled", False)

    if _is_cancelled():
        logging.info(f"[YD-UP] Задача {task_id} отменена — upload пропущен")
        pending = session.get("pending")
        if isinstance(pending, dict):
            await asyncio.shield(_yd_cleanup_task(
                task_id,
                pending.get("session_key"),
                pending.get("task_dir"),
                pending.get("owner_user_id"),
                pending.get("chat_id"),
                pending.get("nonce"),
            ))
        return

    pending = session["pending"]
    if not isinstance(pending, dict):
        return

    prepared = pending["prepared"]
    target_base = pending["target_base"]
    task_dir = pending["task_dir"]
    owner_user_id = pending["owner_user_id"]
    chat_id = pending["chat_id"]
    session_key = pending["session_key"]
    nonce = pending["nonce"]
    quality = pending.get("quality", "2k")

    if status_msg is not None and hasattr(status_msg, "message_id"):
        pending["status_message_id"] = status_msg.message_id

    zip_tmp_dir: Optional[Path] = None
    cleanup_done = False

    logging.info(
        f"[YD-UP] Старт v4.0: task_id={task_id}, файлов={len(prepared)}, "
        f"target_base={target_base!r}"
    )

    try:
        zip_tmp_dir = Path(tempfile.mkdtemp(prefix=f"pptx2png_{task_id}_"))
        logging.debug(f"[YD-UP] Создана временная папка {zip_tmp_dir}")

        total_uploaded_zip = 0
        total_slides_packed = 0
        total_failed = 0
        report_lines = [f"📁 Обработано файлов: <b>{len(prepared)}</b>\n"]
        links_by_folder: dict[str, str] = {}
        meta = _cat_meta()

        for f_idx, item in enumerate(prepared, start=1):
            _touch_task(task_dir)
            _normalize_item_ranges(item)

            if _is_cancelled():
                logging.info(f"[YD-UP] Отмена перед файлом #{f_idx}")
                return

            file_name = item["file_name"]
            file_name_esc = html_module.escape(file_name)
            file_slug = item["file_slug"]

            if item.get("failed_at_stage"):
                stage = item["failed_at_stage"]
                stage_text = {
                    "download":  "ошибка скачивания",
                    "normalize": "ошибка подготовки .ppt → .pptx",
                    "count":     "не удалось определить число слайдов",
                    "convert":   "ошибка конвертации",
                    "no_pngs":   "нет PNG",
                }.get(stage, stage)
                report_lines.append(f"❌ {file_name_esc} — {stage_text}")
                total_failed += 1
                continue

            file_path = item.get("file_path")
            if not file_path or not Path(file_path).exists():
                logging.error(f"[YD-UP] {file_name}: PPTX не найден")
                report_lines.append(f"❌ {file_name_esc} — файл PPTX не найден")
                total_failed += 1
                continue

            selected = set(item.get("selected_categories") or set())
            if not selected:
                logging.warning(
                    f"[YD-UP] {file_name}: ничего не выбрано — пропускаем"
                )
                report_lines.append(
                    f"⏭ {file_name_esc} — ничего не выбрано"
                )
                continue

            selected_labels = []
            for cat in _cat_order():
                if cat in selected:
                    cm = meta.get(cat, {})
                    selected_labels.append(
                        f"{cm.get('emoji','❓')} {cm.get('name',cat)}"
                    )
            if "other" in selected:
                selected_labels.append("📄 Остальные")
            label_line = " + ".join(selected_labels) if selected_labels else "—"

            logging.info(
                f"[YD-UP] Файл #{f_idx}: {file_name!r}, "
                f"selected={sorted(selected)}"
            )

            header_lines = [
                f"🎬 <b>Выбранные категории:</b> {label_line}",
                f"📄 Файл: <code>{file_name_esc}</code>",
                "",
            ]
            header = "\n".join(header_lines)

            base_convert_text = f"{header}⚙️ Конвертирую в PNG"

            await _safe_edit(
                status_msg,
                base_convert_text,
                parse_mode="HTML",
                reply_markup=_get_cancel_keyboard(task_id).as_markup(),
            )

            temp_png_dir = task_dir / f"png_{f_idx}"
            temp_png_dir.mkdir(exist_ok=True)

            try:
                pngs, used_pptx = await _yd_with_spinner(
                    status_msg,
                    task_id,
                    base_convert_text,
                    _yd_run_protected(
                        task_id,
                        convert_all_pngs(
                            Path(file_path), temp_png_dir, quality
                        ),
                    ),
                )
            except Exception as e:
                logging.error(
                    f"[YD-UP] {file_name}: ошибка конвертации: {e}",
                    exc_info=True,
                )
                report_lines.append(
                    f"❌ {file_name_esc} — ошибка конвертации"
                )
                total_failed += 1
                continue

            if not pngs:
                logging.warning(f"[YD-UP] {file_name}: не создано PNG")
                report_lines.append(f"❌ {file_name_esc} — нет PNG")
                total_failed += 1
                continue

            pngs_sorted = sorted(pngs, key=lambda p: p.name)

            if _is_cancelled():
                logging.info(f"[YD-UP] Отмена после конвертации {file_name}")
                return

            pngs_by_cat: Dict[str, List[Path]] = {
                cat: [] for cat in _cat_order()
            }
            pngs_by_cat["other"] = []

            for slide_idx, png_path in enumerate(pngs_sorted, start=1):
                cat = _classify_slide(item, slide_idx)
                if cat is None:
                    pngs_by_cat["other"].append(png_path)
                else:
                    pngs_by_cat[cat].append(png_path)

            logging.debug(
                f"[YD-UP] {file_name}: классификация "
                + ", ".join(
                    f"{c}={len(pngs_by_cat[c])}"
                    for c in list(_cat_order()) + ["other"]
                )
            )

            pptx2png_dir = (
                f"{target_base}/{yandex_state.config.pptx2png_folder}/{file_slug}"
            )

            need_other = (
                "other" in selected and pngs_by_cat["other"]
            )

            if need_other:
                ok0 = await _yd_async_protected(
                    task_id,
                    yandex_state.config.client.ensure_folder(pptx2png_dir),
                )
                if not ok0:
                    logging.error(
                        f"[YD-UP] {file_name}: не удалось создать {pptx2png_dir!r}"
                    )
                    report_lines.append(
                        f"❌ {file_name_esc} — не удалось создать папку остальных"
                    )
                    total_failed += 1
                    continue

            cat_folders_ok = {}
            for cat in _cat_order():
                if cat not in selected or not pngs_by_cat[cat]:
                    continue
                folder = _cat_folder_path(cat, target_base)
                ok = await _yd_async_protected(
                    task_id,
                    yandex_state.config.client.ensure_folder(folder),
                )
                cat_folders_ok[cat] = ok
                if not ok:
                    logging.error(
                        f"[YD-UP] {file_name}: не удалось создать {folder!r}"
                    )

            entry_lines = [f"{f_idx}. 📄 <b>{file_name_esc}</b>"]
            file_uploaded = 0
            file_slides = 0

            for cat in _cat_order():
                if cat not in selected:
                    continue
                pngs_list = pngs_by_cat[cat]
                if not pngs_list:
                    continue
                if not cat_folders_ok.get(cat):
                    entry_lines.append(
                        f"   • ❌ {meta.get(cat,{}).get('emoji','❓')} "
                        f"{meta.get(cat,{}).get('name',cat)}: "
                        f"папка не создана"
                    )
                    total_failed += 1
                    continue

                if _is_cancelled():
                    logging.info(f"[YD-UP] Отмена перед ZIP {cat}")
                    return

                cat_meta = meta.get(cat, {})
                emoji = cat_meta.get("emoji", "❓")
                name = cat_meta.get("name", cat)
                zip_suffix = cat_meta.get("zip_suffix", cat)
                folder_path = _cat_folder_path(cat, target_base)
                cat_data = item.get("categories", {}).get(cat) or {}
                ranges = cat_data.get("ranges") or []
                ranges_text = _format_ranges_text(ranges)

                zip_name = f"{file_slug}_{zip_suffix}.zip"
                zip_path = zip_tmp_dir / zip_name

                try:
                    await _yd_to_thread(
                        task_id,
                        create_zip_stream, pngs_list, zip_path,
                    )
                except Exception as e:
                    logging.error(
                        f"[YD-UP] {file_name}: ZIP {cat} упал: {e}",
                        exc_info=True,
                    )
                    entry_lines.append(
                        f"   • ❌ {emoji} {name}: ошибка упаковки"
                    )
                    total_failed += 1
                    continue

                if _is_cancelled():
                    _safe_unlink(zip_path)
                    return

                zip_size = zip_path.stat().st_size
                remote_path = f"{folder_path}/{zip_name}"

                base_upload_text = (
                    f"{header}"
                    f"📤 Загружаю {emoji} {name} "
                    f"(<code>{file_name_esc}</code>)"
                )

                await _safe_edit(
                    status_msg,
                    base_upload_text,
                    parse_mode="HTML",
                    reply_markup=_get_cancel_keyboard(task_id).as_markup(),
                )

                if _is_cancelled():
                    _safe_unlink(zip_path)
                    return

                ok = await _yd_with_spinner(
                    status_msg,
                    task_id,
                    base_upload_text,
                    _yd_async_protected(
                        task_id,
                        yandex_state.config.client.upload_file(
                            zip_path, remote_path
                        ),
                    ),
                )

                if ok:
                    total_uploaded_zip += 1
                    total_slides_packed += len(pngs_list)
                    file_uploaded += 1
                    file_slides += len(pngs_list)
                    links_by_folder[folder_path] = f"{emoji} {name}"
                    folder_short = folder_path.rsplit("/", 1)[-1]
                    entry_lines.append(
                        f"   • {emoji} {name} ({ranges_text}): "
                        f"{len(pngs_list)} с. → "
                        f"<code>{html_module.escape(folder_short)}/"
                        f"{html_module.escape(zip_name)}</code> "
                        f"({_format_size(zip_size)})"
                    )
                    logging.info(
                        f"[YD-UP] {file_name}: {cat} ZIP загружен "
                        f"({len(pngs_list)} с.)"
                    )
                else:
                    total_failed += 1
                    entry_lines.append(
                        f"   • ❌ {emoji} {name}: не удалось загрузить"
                    )
                    logging.error(f"[YD-UP] {file_name}: {cat} ZIP failed")

                _safe_unlink(zip_path)
                for png in pngs_list:
                    _safe_unlink(png)

            if need_other:
                pngs_list = pngs_by_cat["other"]
                if _is_cancelled():
                    logging.info(f"[YD-UP] Отмена перед ZIP other")
                    return

                zip_name = f"{file_slug}_слайды.zip"
                zip_path = zip_tmp_dir / zip_name

                try:
                    await _yd_to_thread(
                        task_id,
                        create_zip_stream, pngs_list, zip_path,
                    )
                except Exception as e:
                    logging.error(
                        f"[YD-UP] {file_name}: ZIP other упал: {e}",
                        exc_info=True,
                    )
                    entry_lines.append(
                        "   • ❌ 📄 Остальные: ошибка упаковки"
                    )
                    total_failed += 1
                else:
                    if _is_cancelled():
                        _safe_unlink(zip_path)
                        return

                    zip_size = zip_path.stat().st_size
                    remote_path = f"{pptx2png_dir}/{zip_name}"

                    base_upload_text = (
                        f"{header}"
                        f"📤 Загружаю 📄 Остальные "
                        f"(<code>{file_name_esc}</code>)"
                    )

                    await _safe_edit(
                        status_msg,
                        base_upload_text,
                        parse_mode="HTML",
                        reply_markup=_get_cancel_keyboard(task_id).as_markup(),
                    )

                    if _is_cancelled():
                        _safe_unlink(zip_path)
                        return

                    ok = await _yd_with_spinner(
                        status_msg,
                        task_id,
                        base_upload_text,
                        _yd_async_protected(
                            task_id,
                            yandex_state.config.client.upload_file(
                                zip_path, remote_path
                            ),
                        ),
                    )

                    folder_short = (
                        f"{yandex_state.config.pptx2png_folder}/{file_slug}"
                    )

                    if ok:
                        total_uploaded_zip += 1
                        total_slides_packed += len(pngs_list)
                        file_uploaded += 1
                        file_slides += len(pngs_list)
                        links_by_folder[pptx2png_dir] = "📄 Остальные"
                        entry_lines.append(
                            f"   • 📄 Остальные: {len(pngs_list)} с. → "
                            f"<code>{html_module.escape(folder_short)}/"
                            f"{html_module.escape(zip_name)}</code> "
                            f"({_format_size(zip_size)})"
                        )
                        logging.info(
                            f"[YD-UP] {file_name}: other ZIP загружен "
                            f"({len(pngs_list)} с.)"
                        )
                    else:
                        total_failed += 1
                        entry_lines.append(
                            "   • ❌ 📄 Остальные: не удалось загрузить"
                        )

                    _safe_unlink(zip_path)
                    for png in pngs_list:
                        _safe_unlink(png)

            if file_uploaded == 0:
                entry_lines.append("   • ℹ️ Ничего не загружено")

            incomplete_warning = item.get("incomplete_warning")
            if incomplete_warning:
                entry_lines.append(f"   • {incomplete_warning}")

            report_lines.append("\n".join(entry_lines))

        if _is_cancelled():
            logging.info(f"[YD-UP] Отмена перед отправкой отчёта")
            return

        if total_failed > 0:
            report_lines.append(
                f"\n⚠️ Всего загружено архивов: <b>{total_uploaded_zip}</b>\n"
                f"📊 Всего слайдов: <b>{total_slides_packed}</b>\n"
                f"❌ Ошибок: <b>{total_failed}</b>"
            )
        else:
            report_lines.append(
                f"\n📊 Всего загружено архивов: <b>{total_uploaded_zip}</b>\n"
                f"📊 Всего слайдов: <b>{total_slides_packed}</b>"
            )

        if links_by_folder:
            report_lines.append("\n🔗 <b>Ссылки на Яндекс.Диск:</b>")
            for folder_path, label in links_by_folder.items():
                url = _yd_public_url(folder_path)
                folder_name = folder_path.rsplit("/", 1)[-1]
                report_lines.append(
                    f'   • {label}: <a href="{url}">'
                    f'{html_module.escape(folder_name)}/</a>'
                )

        try:
            await status_msg.edit_reply_markup(reply_markup=None)
        except TelegramBadRequest:
            pass
        except Exception:
            pass

        await _yd_send_report(
            bot=bot,
            chat_id=chat_id,
            status_msg=status_msg,
            report_lines=report_lines,
            total_uploaded=total_uploaded_zip,
            total_failed=total_failed,
        )

        logging.info(
            f"[YD-UP] Итог v4.0: zip={total_uploaded_zip}, "
            f"slides={total_slides_packed}, failed={total_failed}"
        )

    except asyncio.CancelledError:
        logging.info(f"[YD-UP] Worker {task_id} отменён — выполняю cleanup")
        cleanup_done = True
        try:
            await asyncio.shield(_yd_cleanup_task(
                task_id, session_key, task_dir,
                owner_user_id, chat_id, nonce,
            ))
        except Exception as e:
            logging.error(
                f"[YD-UP] cleanup после отмены упал: {e}", exc_info=True
            )
        raise

    except Exception as e:
        logging.error(f"[YD-UP] Ошибка: {e}", exc_info=True)
        cleanup_done = True
        await _yd_cleanup_task(
            task_id, session_key, task_dir,
            owner_user_id, chat_id, nonce,
            bot=bot, status_msg=status_msg, error=e,
        )
        return
    finally:
        if zip_tmp_dir is not None and zip_tmp_dir.exists():
            try:
                shutil.rmtree(zip_tmp_dir)
                logging.debug(f"[YD-UP] Удалена временная папка {zip_tmp_dir}")
            except Exception as e:
                logging.warning(
                    f"[YD-UP] Не удалось удалить {zip_tmp_dir}: {e}",
                    exc_info=True,
                )

        if not cleanup_done:
            await _yd_cleanup_task(
                task_id, session_key, task_dir,
                owner_user_id, chat_id, nonce,
            )


async def _yd_convert_and_upload(
    bot: Bot,
    task_id: str,
    status_msg,
) -> asyncio.Task:
    """Запускает _yd_convert_and_upload_impl как detached worker."""
    return await _run_worker_detached(
        task_id,
        lambda: _yd_convert_and_upload_impl(
            bot=bot, task_id=task_id, status_msg=status_msg,
        ),
        name="convert",
    )


async def _yd_send_report(
    bot: Bot,
    chat_id: int,
    status_msg,
    report_lines: list,
    total_uploaded: int,
    total_failed: int,
):
    """Отправляет отчёт пользователю (по частям, если длинный)."""
    MAX_MSG_LEN = 3500
    chunks = []
    current_chunk = []
    current_len = 0

    for line in report_lines:
        line_len = len(line) + 1
        if current_len + line_len > MAX_MSG_LEN and current_chunk:
            chunks.append("\n".join(current_chunk))
            current_chunk = [line]
            current_len = line_len
        else:
            current_chunk.append(line)
            current_len += line_len

    if current_chunk:
        chunks.append("\n".join(current_chunk))

    delivered = 0
    failed_chunks = []
    first_delivered = False

    if chunks:
        try:
            await status_msg.edit_text(
                chunks[0], parse_mode="HTML", disable_web_page_preview=True
            )
            delivered += 1
            first_delivered = True
        except Exception as e:
            logging.error(
                f"Ошибка edit_text первой части: {e}", exc_info=True
            )
            try:
                await bot.send_message(
                    chat_id=chat_id,
                    text=chunks[0],
                    parse_mode="HTML",
                    disable_web_page_preview=True,
                )
                delivered += 1
                first_delivered = True
            except Exception as e2:
                logging.error(
                    f"Не удалось отправить первую часть: {e2}", exc_info=True
                )
                failed_chunks.append(1)

    if not first_delivered:
        try:
            fallback = (
                f"⚠️ Не удалось показать полный отчёт.\n"
                f"📊 Загружено: {total_uploaded}"
            )
            if total_failed:
                fallback += f"\n❌ Ошибок: {total_failed}"
            await bot.send_message(chat_id=chat_id, text=fallback)
        except Exception as e:
            logging.error(
                f"Не удалось отправить fallback-отчёт: {e}", exc_info=True
            )

    for i, chunk in enumerate(chunks[1:], start=2):
        try:
            await bot.send_message(
                chat_id=chat_id,
                text=chunk,
                parse_mode="HTML",
                disable_web_page_preview=True,
            )
            delivered += 1
        except Exception as e:
            logging.error(f"Ошибка отправки части {i}: {e}", exc_info=True)
            failed_chunks.append(i)

    if failed_chunks and delivered > 0:
        try:
            await bot.send_message(
                chat_id=chat_id,
                text=(
                    f"⚠️ Не удалось доставить {len(failed_chunks)} "
                    f"из {len(chunks)} частей отчёта. Проверьте Яндекс.Диск."
                ),
            )
        except Exception as e:
            logging.error(f"Не удалось отправить предупреждение: {e}")
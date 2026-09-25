# ==========================================
# converter_engine.py — КОНВЕРТЕР PPTX → PNG (v1.5)
# ==========================================
# Изменения v1.5:
#   • Добавлена count_slides_via_libreoffice() — fallback для подсчёта
#     числа слайдов через LibreOffice + PyMuPDF, когда python-pptx
#     не может открыть файл.
# ==========================================

import os
import shutil
import subprocess
import zipfile
import re
import asyncio
import logging
from pathlib import Path
from typing import List, Tuple, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
import fitz


# ==========================================
# ВСПОМОГАТЕЛЬНЫЕ КЛАССЫ
# ==========================================

class FakeArgs:
    def __init__(self, quality, keep_pdf, dark_mode=True, zip_mode=True,
                 clean=True, output_dir=None):
        self.quality = quality
        self.keep_pdf = keep_pdf
        self.dark_mode = dark_mode
        self.zip = zip_mode
        self.clean = clean
        self.output_dir = output_dir


# ==========================================
# DARK MODE
# ==========================================

def make_dark_mode(pptx_path, temp_output_path):
    """
    Создаёт копию презентации с чёрным фоном, белым текстом
    и без фоновых картинок (>80% площади слайда).
    """
    prs = Presentation(pptx_path)
    BLACK = RGBColor(0, 0, 0)
    WHITE = RGBColor(255, 255, 255)

    for slide in prs.slides:
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BLACK

        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                slide_area = prs.slide_width * prs.slide_height
                shape_area = shape.width * shape.height
                if shape_area / slide_area > 0.8:
                    sp = shape._element
                    sp.getparent().remove(sp)
                    continue

            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = WHITE

            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.font.color.rgb = WHITE

    prs.save(temp_output_path)


# ==========================================
# РАЗРЕШЕНИЕ
# ==========================================

def get_resolution_multiplier(quality_str, page):
    rect = page.rect
    long_side = max(rect.width, rect.height)
    if quality_str == "2k":
        target = 2560
    elif quality_str == "4k":
        target = 3840
    else:
        return 2.0
    return target / long_side


# ==========================================
# КОНВЕРТАЦИЯ PPT/PPTX → PDF
# ==========================================

def _find_libreoffice() -> str:
    """Ищет LibreOffice в системе."""
    mac_path = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if os.path.exists(mac_path):
        return mac_path
    if shutil.which("soffice") is not None:
        return "soffice"
    raise FileNotFoundError(
        "LibreOffice не найден! Установите: sudo apt install libreoffice"
    )


def pptx_to_pdf_crossplatform(pptx_path: Path, output_dir: Path) -> Path:
    """Конвертирует PPTX в PDF через LibreOffice."""
    libreoffice_path = _find_libreoffice()

    cmd = [
        libreoffice_path,
        "--headless",
        "--convert-to", "pdf",
        "--outdir", str(output_dir),
        str(pptx_path),
    ]
    subprocess.run(
        cmd,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        check=True,
    )
    return output_dir / f"{pptx_path.stem}.pdf"


def ppt_to_pptx_crossplatform(ppt_path: Path, output_dir: Path) -> Path:
    """Конвертирует PPT (старый формат) в PPTX через LibreOffice."""
    libreoffice_path = _find_libreoffice()

    cmd = [
        libreoffice_path,
        "--headless",
        "--convert-to", "pptx",
        "--outdir", str(output_dir),
        str(ppt_path),
    ]
    subprocess.run(
        cmd,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        check=True,
    )
    return output_dir / f"{ppt_path.stem}.pptx"


def count_slides_via_libreoffice(
    pptx_path: Path, work_dir: Path
) -> Optional[int]:
    """
    Fallback-способ узнать число слайдов: LibreOffice → PDF → fitz.

    Используется, когда python-pptx не может открыть файл (например,
    специфичный OOXML), но LibreOffice рендерит его нормально.

    Возвращает количество слайдов или None при любой ошибке.
    PDF удаляется сразу после подсчёта.
    """
    try:
        pdf_path = pptx_to_pdf_crossplatform(pptx_path, work_dir)
        if not pdf_path or not pdf_path.exists():
            return None
        doc = fitz.open(pdf_path)
        count = len(doc)
        doc.close()
        try:
            pdf_path.unlink()
        except Exception:
            pass
        return count if count > 0 else None
    except Exception as e:
        logging.warning(
            f"count_slides_via_libreoffice({pptx_path}): {e}"
        )
        return None


# ==========================================
# PDF → PNG
# ==========================================

def pdf_to_png_fast(
    pdf_path: Path, output_dir: Path, quality: str
) -> Tuple[int, List[Path]]:
    """
    Конвертирует PDF в набор PNG (по слайду на страницу).
    Возвращает (total_pages, [пути к PNG]).
    """
    doc = fitz.open(pdf_path)
    total_pages = len(doc)
    z_fill_len = max(2, len(str(total_pages)))

    created_files: List[Path] = []
    for page_num in range(total_pages):
        page = doc.load_page(page_num)
        zoom = get_resolution_multiplier(quality, page)
        mat = fitz.Matrix(zoom, zoom)

        pix = page.get_pixmap(matrix=mat)
        slide_index = str(page_num + 1).zfill(z_fill_len)
        png_path = output_dir / f"slide_{slide_index}.png"
        pix.save(str(png_path))
        created_files.append(png_path)

    doc.close()
    return total_pages, created_files


# ==========================================
# ZIP
# ==========================================

def create_zip_stream(
    file_paths: List[Path], output_path: Path, compress_level: int = 6
) -> Path:
    """
    Создаёт ZIP-архив из списка файлов.
    Сохраняет имена файлов без путей (arcname=f.name).
    """
    with zipfile.ZipFile(
        output_path, 'w', zipfile.ZIP_DEFLATED, compresslevel=compress_level
    ) as zf:
        for fpath in file_paths:
            if fpath.exists():
                zf.write(fpath, arcname=fpath.name)
    return output_path


# ==========================================
# ZIP EXTRACT
# ==========================================

def extract_zip_if_needed(zip_path: Path, extract_dir: Path) -> Optional[Path]:
    """
    Распаковывает ZIP и возвращает путь к первой найденной презентации.
    Игнорирует временные файлы ~$.
    """
    with zipfile.ZipFile(zip_path, 'r') as zip_ref:
        zip_ref.extractall(extract_dir)

    for ext in ("*.pptx", "*.PPTX", "*.ppt", "*.PPT"):
        for file in extract_dir.glob(ext):
            if not file.name.startswith("~$"):
                return file
    return None


# ==========================================
# URL → DIRECT DOWNLOAD
# ==========================================

def convert_to_direct_download(url: str) -> str:
    """Преобразует ссылки Google Docs/Drive в прямые ссылки на скачивание."""
    url = url.strip()

    if "docs.google.com/presentation" in url:
        match = re.search(r'/presentation/d/([a-zA-Z0-9_-]+)', url)
        if match:
            return (
                f"https://docs.google.com/presentation/d/"
                f"{match.group(1)}/export/pptx"
            )

    if "://google.com" in url or "drive.google.com" in url:
        match = re.search(r'/file/d/([a-zA-Z0-9_-]+)', url)
        if match:
            return f"https://drive.google.com/uc?export=download&id={match.group(1)}"

    match_id = re.search(r'[?&]id=([a-zA-Z0-9_-]+)', url)
    if match_id:
        return f"https://drive.google.com/uc?export=download&id={match_id.group(1)}"

    return url


# ==========================================
# ВЫСОКОУРОВНЕВАЯ ОБЁРТКА (используется в handlers.py)
# ==========================================

async def convert_all_pngs(
    pptx_path: Path,
    output_dir: Path,
    quality: str,
) -> Tuple[List[Path], Path]:
    """
    Конвертирует PPTX (или PPT) в список PNG.
    Возвращает (список PNG, путь к .pptx, использованному для рендера).

    Для .ppt сначала конвертирует в .pptx. Возвращённый `used_pptx`
    НЕ удаляется — он нужен для извлечения заметок вызывающим кодом.
    Вызывающий код должен сам удалить его после использования.

    Запускается в отдельном потоке (asyncio.to_thread), чтобы не блокировать event loop.
    """
    def _sync_convert() -> Tuple[List[Path], Path]:
        # 1. PPT → PPTX
        if pptx_path.suffix.lower() == '.ppt':
            pptx_converted = ppt_to_pptx_crossplatform(pptx_path, output_dir)
        else:
            pptx_converted = pptx_path

        # 2. Dark mode
        temp_dark_pptx = output_dir / f"temp_dark_{pptx_converted.name}"
        make_dark_mode(pptx_converted, temp_dark_pptx)

        # 3. PDF
        pdf_path = pptx_to_pdf_crossplatform(temp_dark_pptx, output_dir)

        # 4. PNG
        total_slides, png_paths = pdf_to_png_fast(pdf_path, output_dir, quality)

        # 5. Очистка промежуточных файлов
        if pdf_path.exists():
            try:
                pdf_path.unlink()
            except Exception:
                pass
        if temp_dark_pptx.exists():
            try:
                temp_dark_pptx.unlink()
            except Exception:
                pass

        # ❌ НЕ удаляем pptx_converted здесь — он нужен для заметок
        return png_paths, pptx_converted

    pngs, used_pptx = await asyncio.to_thread(_sync_convert)
    return pngs, used_pptx


# ==========================================
# LEGACY: process_file_local (используется CLI)
# ==========================================

def process_file_local(pptx_path, args):
    """
    Обрабатывает один PPTX-файл: dark mode → PDF → PNG → ZIP.
    Используется CLI convert.py. В боте не вызывается.
    """
    file_output_dir = Path(args.output_dir) / f"{pptx_path.stem}_output"
    file_output_dir.mkdir(parents=True, exist_ok=True)

    current_pptx = pptx_path
    temp_dark_pptx = None
    zip_path = None

    try:
        if args.dark_mode:
            temp_dark_pptx = file_output_dir / f"temp_dark_{pptx_path.name}"
            make_dark_mode(pptx_path, temp_dark_pptx)
            current_pptx = temp_dark_pptx

        pdf_path = pptx_to_pdf_crossplatform(current_pptx, file_output_dir)
        total_slides, generated_pngs = pdf_to_png_fast(
            pdf_path, file_output_dir, args.quality
        )

        if args.zip:
            zip_path = Path(args.output_dir) / f"{pptx_path.stem}_output.zip"
            create_zip_stream(generated_pngs, zip_path)

        if not args.keep_pdf and pdf_path.exists():
            pdf_path.unlink()
        if temp_dark_pptx and temp_dark_pptx.exists():
            temp_dark_pptx.unlink()

        return zip_path

    except Exception as e:
        if temp_dark_pptx and temp_dark_pptx.exists():
            temp_dark_pptx.unlink()
        raise e
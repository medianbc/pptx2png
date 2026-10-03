"""Utilities shared by the bot and local presentation workflow."""

import logging
from typing import Dict, Tuple

from pptx import Presentation


def extract_speaker_notes(file_path: str) -> Tuple[bool, Dict[int, str], bool]:
    """Extract notes by slide number and indicate whether extraction was partial."""
    try:
        presentation = Presentation(file_path)
        notes = {}
        incomplete = False

        for slide_number, slide in enumerate(presentation.slides, start=1):
            if not slide.has_notes_slide:
                continue
            try:
                text = slide.notes_slide.notes_text_frame.text or ""
                if text.strip():
                    notes[slide_number] = text
            except Exception as error:
                logging.warning(
                    "Не удалось прочитать заметки слайда %s: %s",
                    slide_number,
                    error,
                )
                incomplete = True

        if incomplete:
            logging.warning(
                "Заметки извлечены частично (%s слайдов), некоторые слайды не прочитаны",
                len(notes),
            )
        else:
            logging.info(
                "Извлечены заметки для %s слайдов из %s",
                len(notes),
                file_path,
            )

        return True, notes, incomplete
    except Exception as error:
        logging.error("Ошибка извлечения заметок: %s", error, exc_info=True)
        return False, {}, False

# ==========================================
# yandex_flow.py — ОРКЕСТРАЦИЯ ЯНДЕКС.ДИСКА (v3.4)
# ==========================================
# Изменения v3.4:
#   • Убран дублированный [YD-PREP] Старт в логах
#   • Единая утилита _picker_is_active — все три проверки
#     (cmd_sunday верхняя, cmd_sunday защитная, yd_pick)
#     используют одну семантику «активная сессия»
#   • session_key сохраняется в sessions[task_id] при регистрации
#   • yd_task_cancel_callback немедленно убирает task_id из
#     picker["task_ids"] — отменённая задача не блокирует новый /sunday
#   • _yd_cleanup_task мутирует picker только при совпадении nonce
# ==========================================
# Изменения v3.3:
#   • Race-condition fix: /sunday не плодит параллельные сессии
#   • Сверка nonce при регистрации task_id
#   • yd_pick: try/except + сброс processing
# ==========================================
# Изменения v3.2:
#   • _count_slides_in_ranges, опечатка prepared, короткие task_id/nonce,
#     шапка режима, спиннер, нормализация .ppt, кэш total_slides,
#     уникальная подпапка, skip-режим
# ==========================================

import asyncio
import html as html_module
import logging
import os
import secrets
import shutil
import tempfile
import time
import urllib.parse
from pathlib import Path
from typing import Optional

from aiogram import Router, F, types, Bot
from aiogram.filters import Command
from aiogram.types import InlineKeyboardButton
from aiogram.utils.keyboard import InlineKeyboardBuilder

import yandex_state
from yandex_state import (
    sessions,
    yd_session_lock,
    yd_active_tasks,
    yd_try_acquire,
    yd_release,
    yd_is_active,
)

from yandex_disk import (
    YandexDiskError,
    get_nearest_sunday,
    month_folder_name,
    resolve_sunday_paths,
    find_pptx_in_source,
)
from structure import safe_folder_name
from sermon_detector import find_sermon_range
from utils import extract_speaker_notes
from converter_engine import (
    convert_all_pngs,
    create_zip_stream,
    ppt_to_pptx_crossplatform,
    librenormalize_to_pptx,
)


router = Router()


# ==========================================
# УТИЛИТЫ
# ==========================================

def _touch_task(task_dir: Path):
    """Обновляет mtime папки — чтобы cleaner не удалял активную задачу."""
    if task_dir and task_dir.exists():
        try:
            os.utime(task_dir, None)
        except Exception as e:
            logging.error(f"Ошибка touch для {task_dir}: {e}")


def _safe_delete_task_dir(task_dir: Path):
    """Безопасно удаляет папку задачи."""
    if task_dir and task_dir.exists():
        try:
            shutil.rmtree(task_dir)
            logging.info(f"🧹 Удалена папка задачи: {task_dir}")
        except Exception as e:
            logging.error(f"Ошибка удаления папки {task_dir}: {e}")


def _safe_unlink(path: Path):
    """Безопасно удаляет файл, логирует ошибку при неудаче."""
    if path is None:
        return
    try:
        if path.exists():
            path.unlink()
    except Exception as e:
        logging.warning(f"Не удалось удалить {path}: {e}")


async def _safe_edit(msg, text: str, **kwargs) -> bool:
    """Безопасный edit_text с логированием ошибок."""
    if msg is None:
        return False
    try:
        await msg.edit_text(text, **kwargs)
        return True
    except Exception as e:
        logging.warning(f"Не удалось обновить статус-сообщение: {e}", exc_info=True)
        return False


def _format_ranges_text(ranges, start=None, end=None) -> str:
    """Форматирует список диапазонов в '1–2, 8–9'."""
    if ranges:
        return ", ".join(
            f"{s}–{e}" if s != e else str(s) for s, e in ranges
        )
    if start is not None and end is not None:
        if start != end:
            return f"{start}–{end}"
        return str(start)
    return "—"


def _is_sermon_slide(item: dict, slide_idx: int) -> bool:
    """Проверяет, относится ли слайд к проповеди (по ranges или start..end)."""
    ranges = item.get("ranges")
    if ranges:
        return any(s <= slide_idx <= e for s, e in ranges)
    start = item.get("start")
    end = item.get("end")
    return start is not None and start <= slide_idx <= end


def _yd_public_url(disk_path: str) -> str:
    """Строит корректный кликабельный URL Яндекс.Диска."""
    path = disk_path.lstrip("/")
    encoded = urllib.parse.quote(path, safe="/")
    return f"https://disk.yandex.ru/client/disk/{encoded}"


def _format_size(num_bytes: int) -> str:
    """Человекочитаемый размер файла."""
    if num_bytes < 1024:
        return f"{num_bytes} Б"
    if num_bytes < 1024 * 1024:
        return f"{num_bytes / 1024:.1f} КБ"
    if num_bytes < 1024 * 1024 * 1024:
        return f"{num_bytes / (1024 * 1024):.1f} МБ"
    return f"{num_bytes / (1024 * 1024 * 1024):.2f} ГБ"


def _get_cancel_keyboard(task_id: str) -> InlineKeyboardBuilder:
    """Клавиатура с единственной кнопкой «Отменить задачу»."""
    kb = InlineKeyboardBuilder()
    kb.row(
        InlineKeyboardButton(
            text="❌ Отменить задачу",
            callback_data=f"yd_task_cancel:{task_id}",
        ),
    )
    return kb


def _count_slides_for_range(start, end, total):
    """Возвращает количество слайдов в диапазоне [start, end]."""
    if start is None or end is None:
        return 0
    return max(0, min(end, total) - max(start, 1) + 1)


def _count_slides_in_ranges(ranges, total_slides: int) -> int:
    """
    Считает количество УНИКАЛЬНЫХ слайдов, попадающих в объединение диапазонов.

    ranges: список кортежей (start, end) или None
    total_slides: общее количество слайдов (для клиппинга)

    Пример:
        ranges=[(1,1),(10,10)], total=20 → 2
        ranges=[(1,5),(3,8)],   total=20 → 8  (пересечение не двоится)
    """
    if not ranges or total_slides <= 0:
        return 0

    unique_slides = set()
    for s, e in ranges:
        lo = max(1, s)
        hi = min(total_slides, e)
        if lo > hi:
            continue
        unique_slides.update(range(lo, hi + 1))
    return len(unique_slides)


def _picker_is_active(picker: Optional[dict]) -> bool:
    """
    Пикер считается активным, если:
      • processing=True (yd_pick выбрал файл, но _yd_prepare_files
        ещё не зарегистрировал task_id), ИЛИ
      • есть живые (не отменённые) задачи в task_ids.

    Отменённые задачи (session["cancelled"]=True) игнорируются:
    cleanup доведёт их до конца в фоне, но новый /sunday они
    блокировать не должны.

    ВАЖНО: все проверки активной сессии в cmd_sunday / yd_pick
    должны использовать эту функцию — единая семантика.
    """
    if not picker:
        return False
    if picker.get("processing"):
        return True
    for tid in picker.get("task_ids", []):
        task_sess = sessions.get(tid)
        if task_sess is not None and not task_sess.get("cancelled"):
            return True
    return False


# ==========================================
# СПИННЕР ПРОГРЕССА
# ==========================================

_MODE_LABELS = {
    "sermon": "🎯 Только проповедь",
    "other":  "📄 Только остальные",
    "both":   "📦 Проповедь + остальные",
}

# Реестр активных спиннеров: task_id -> (stop_event, spinner_task)
_active_spinners: dict[str, tuple[asyncio.Event, asyncio.Task]] = {}

# Реестр in-flight операций задачи (thread/async ops, которые НЕ
# должны обрываться на полпути). task_id -> set[asyncio.Task]
_active_ops: dict[str, set[asyncio.Task]] = {}


async def _yd_run_protected(task_id: str, coro):
    """
    Запускает корутину как отдельный Task, регистрирует в _active_ops[task_id]
    и защищает от внешней отмены через asyncio.shield.

    Зачем: при worker.cancel() обёртка умирает, но underlying-операция
    (особенно executor-поток внутри asyncio.to_thread) продолжает жить.
    Cleanup должен дождаться её завершения ПЕРЕД удалением task_dir.
    """
    async def _runner():
        try:
            return await coro
        finally:
            bucket = _active_ops.get(task_id)
            if bucket is not None:
                bucket.discard(asyncio.current_task())
                if not bucket:
                    _active_ops.pop(task_id, None)

    op_task = asyncio.create_task(_runner())
    _active_ops.setdefault(task_id, set()).add(op_task)
    # shield: внешняя отмена не убьёт op_task — cleanup дождётся его.
    return await asyncio.shield(op_task)


async def _yd_to_thread(task_id: str, fn, *args, **kwargs):
    """asyncio.to_thread + регистрация в _active_ops."""
    return await _yd_run_protected(
        task_id, asyncio.to_thread(fn, *args, **kwargs)
    )


async def _yd_deferred_task_dir_cleanup(task_dir: Path, ops: list) -> None:
    """
    Фоновое удаление task_dir: ждём завершения in-flight ops, потом rmtree.
    """
    try:
        await asyncio.gather(*ops, return_exceptions=True)
    except Exception as e:
        logging.debug(f"[YD-DEFERRED] gather: {e}")
    try:
        if task_dir and task_dir.exists():
            shutil.rmtree(task_dir)
            logging.info(f"🧹 [YD-DEFERRED] Удалена папка {task_dir} после ops")
    except Exception as e:
        logging.error(f"[YD-DEFERRED] ошибка удаления {task_dir}: {e}")

def _mode_label(mode: str) -> str:
    """Человекочитаемая метка выбранного режима конвертации."""
    if mode == "skip":
        return "⏭ Пропущено"
    return _MODE_LABELS.get(mode, f"❓ {mode}")


async def _yd_progress_spinner(
    status_msg,
    task_id: str,
    base_text: str,
    stop_event: asyncio.Event,
    interval: float = 1.2,
) -> None:
    """Циклически дописывает '.', '..', '...' в конец статусного сообщения."""
    dots = [".", "..", "..."]
    idx = 0
    while not stop_event.is_set():
        try:
            await asyncio.wait_for(stop_event.wait(), timeout=interval)
            return  # stop_event установлен
        except asyncio.TimeoutError:
            pass

        if stop_event.is_set():
            return

        try:
            await status_msg.edit_text(
                f"{base_text} {dots[idx]}",
                parse_mode="HTML",
                reply_markup=_get_cancel_keyboard(task_id).as_markup(),
            )
        except asyncio.CancelledError:
            raise
        except Exception as e:
            # Telegram часто отдаёт "message is not modified" — это не ошибка
            logging.debug(f"[YD-SPINNER] edit_text: {e}")

        idx = (idx + 1) % len(dots)


async def _yd_stop_spinner(task_id: str, wait_timeout: float = 2.0) -> None:
    """
    Останавливает активный спиннер задачи (если есть) и дожидается его выхода.

    Вызывается:
      - в finally у _yd_with_spinner (когда корутина сама завершилась);
      - в yd_task_cancel_callback (до редактирования сообщения).
    """
    entry = _active_spinners.pop(task_id, None)
    if entry is None:
        return
    stop_event, spinner_task = entry
    stop_event.set()
    try:
        await asyncio.wait_for(spinner_task, timeout=wait_timeout)
    except (asyncio.TimeoutError, asyncio.CancelledError):
        spinner_task.cancel()
        try:
            await spinner_task
        except (asyncio.CancelledError, Exception):
            pass


async def _yd_with_spinner(status_msg, task_id: str, base_text: str, coro):
    """
    Запускает корутину `coro`, параллельно анимируя статусное сообщение.
    Гарантированно останавливает спиннер и возвращает результат coro.
    """
    stop_event = asyncio.Event()
    spinner_task = asyncio.create_task(
        _yd_progress_spinner(status_msg, task_id, base_text, stop_event)
    )
    _active_spinners[task_id] = (stop_event, spinner_task)
    try:
        return await coro
    finally:
        entry = _active_spinners.get(task_id)
        if entry is not None and entry[1] is spinner_task:
            _active_spinners.pop(task_id, None)
        stop_event.set()
        try:
            await asyncio.wait_for(spinner_task, timeout=2.0)
        except (asyncio.TimeoutError, asyncio.CancelledError):
            spinner_task.cancel()
            try:
                await spinner_task
            except (asyncio.CancelledError, Exception):
                pass


# ==========================================
# ПАЙПЛАЙН ПОДГОТОВКИ (без конвертации)
# ==========================================

async def _yd_prepare_files(
    callback: types.CallbackQuery,
    bot: Bot,
    SHM_DIR: str,
    user_mgr,
    files_to_process: list,
    sunday,
    paths: dict,
    session_key: str,
    nonce: str,
):
    """
    Первая фаза: скачивание + извлечение заметок + определение проповеди.
    БЕЗ конвертации в PNG — она будет позже, после подтверждения режима.

    Сохраняет результат в sessions[task_id]["pending"].
    """
    status_msg = callback.message
    owner_user_id = callback.from_user.id
    chat_id = callback.message.chat.id

    # ✅ Укороченный task_id — не дублирует user_id (он есть в sessions).
    # Укладывается в лимит callback_data 64 байта вместе с nonce и mode.
    task_id = f"yd_task_{secrets.token_hex(6)}"
    task_dir = Path(SHM_DIR) / task_id

    # ✅ Ранняя регистрация. Под локом — только in-memory операции:
    # никаких await на Telegram, никаких файловых операций.
    registration_status = "ok"  # "ok" | "no_picker" | "wrong_nonce"

    async with yd_session_lock:
        picker = sessions.get(session_key)
        if picker is None:
            registration_status = "no_picker"
        elif picker.get("nonce") != nonce:
            logging.warning(
                f"[YD-PREP] picker.nonce={picker.get('nonce')!r} ≠ "
                f"expected nonce={nonce!r} для session_key={session_key!r} — "
                f"сессия заменена параллельной командой, прерываем"
            )
            registration_status = "wrong_nonce"
        else:
            # Регистрируем задачу сразу — source of truth.
            picker.setdefault("task_ids", []).append(task_id)
            picker["processing"] = False
            sessions[task_id] = {
                "user_id": owner_user_id,
                "chat_id": chat_id,
                "session_key": session_key,  # нужно для отмены на ранних стадиях
                "pending": None,
                "nonce": nonce,
                "created_at": time.time(),
                "cancelled": False,
            }
            yd_active_tasks.add(task_id)

    if registration_status == "no_picker":
        await _safe_edit(status_msg, "❌ Сессия была отменена.")
        return

    if registration_status == "wrong_nonce":
        await _safe_edit(
            status_msg,
            "❌ Сессия была заменена другой командой.\n"
            "Запустите /sunday заново."
        )
        return

    # ✅ mkdir ВНЕ лока. Если упадёт — снимаем флаг уже в except ниже.
    try:
        task_dir.mkdir(parents=True, exist_ok=True)
    except Exception as e:
        logging.error(
            f"[YD-PREP] mkdir упал для {task_id}: {e}",
            exc_info=True,
        )
        # Откатываем регистрацию: убираем task_id из сессий.
        try:
            async with yd_session_lock:
                picker = sessions.get(session_key)
                if picker is not None:
                    task_ids = picker.get("task_ids")
                    if isinstance(task_ids, list) and task_id in task_ids:
                        task_ids.remove(task_id)
                sessions.pop(task_id, None)
                yd_active_tasks.discard(task_id)
        except Exception as cleanup_err:
            logging.error(
                f"[YD-PREP] не удалось откатить регистрацию: {cleanup_err}",
                exc_info=True,
            )
        _safe_delete_task_dir(task_dir)
        await _safe_edit(
            status_msg,
            f"❌ Не удалось создать папку задачи: "
            f"<code>{html_module.escape(str(e)[:200])}</code>",
        )
        return

    status_message_id = status_msg.message_id if status_msg is not None else None

    logging.info(
        f"[YD-PREP] Старт: task_id={task_id}, файлов={len(files_to_process)}"
    )

    try:
        target_base = paths["target"]
        quality = user_mgr.get_user_config(owner_user_id)["quality"]

        prepared = []

        for f_idx, pptx_item in enumerate(files_to_process, start=1):
            _touch_task(task_dir)

            task_sess = sessions.get(task_id)
            if task_sess is None or task_sess.get("cancelled"):
                logging.info(f"[YD-PREP] Задача {task_id} отменена — прерываем подготовку")
                await _yd_cleanup_task(
                    task_id, session_key, task_dir,
                    owner_user_id, chat_id, nonce,
                )
                return

            file_name = pptx_item["name"]
            file_name_esc = html_module.escape(file_name)
            file_size = pptx_item.get("size", 0)

            logging.info(
                f"[YD-PREP] Файл #{f_idx}/{len(files_to_process)}: "
                f"{file_name!r} ({_format_size(file_size)})"
            )

            # ✅ Уникальная подпапка на файл — чтобы .ppt и .pptx с одинаковым
            # stem не перекрывали друг друга в общей task_dir.
            per_file_dir = task_dir / f"src_{f_idx}"
            per_file_dir.mkdir(exist_ok=True)

            # Скачивание с кнопкой отмены + спиннером
            base_dl_text = f"📥 Скачиваю <code>{file_name_esc}</code>"
            await _safe_edit(
                status_msg,
                base_dl_text,
                parse_mode="HTML",
                reply_markup=_get_cancel_keyboard(task_id).as_markup(),
            )

            local_pptx = per_file_dir / file_name
            ok = await _yd_with_spinner(
                status_msg,
                task_id,
                base_dl_text,
                yandex_state.config.client.download_file(
                    pptx_item["path"], local_pptx
                ),
            )
            if not ok:
                logging.error(f"[YD-PREP] {file_name}: скачивание провалилось")
                prepared.append({
                    "file_name": file_name,
                    "file_slug": safe_folder_name(file_name),
                    "failed_at_stage": "download",
                })
                continue

            logging.debug(
                f"[YD-PREP] {file_name}: скачан, "
                f"размер={_format_size(local_pptx.stat().st_size)}"
            )

            # Проверяем отмену после скачивания
            task_sess = sessions.get(task_id)
            if task_sess is None or task_sess.get("cancelled"):
                logging.info(
                    f"[YD-PREP] Задача {task_id} отменена после скачивания "
                    f"{file_name} — прерываем"
                )
                await _yd_cleanup_task(
                    task_id, session_key, task_dir,
                    owner_user_id, chat_id, nonce,
                )
                return

            # ✅ Нормализация .ppt → .pptx (старые форматы).
            if local_pptx.suffix.lower() == ".ppt":
                base_norm_text = f"🔧 Готовлю .ppt → .pptx: <code>{file_name_esc}</code>"
                await _safe_edit(
                    status_msg,
                    base_norm_text,
                    parse_mode="HTML",
                    reply_markup=_get_cancel_keyboard(task_id).as_markup(),
                )
                try:
                    normalized_pptx = await _yd_with_spinner(
                        status_msg,
                        task_id,
                        base_norm_text,
                        _yd_to_thread(
                            task_id,
                            ppt_to_pptx_crossplatform, local_pptx, per_file_dir,
                        ),
                    )
                except Exception as e:
                    logging.error(
                        f"[YD-PREP] {file_name}: .ppt→.pptx упал: {e}",
                        exc_info=True,
                    )
                    prepared.append({
                        "file_name": file_name,
                        "file_slug": safe_folder_name(file_name),
                        "failed_at_stage": "normalize",
                    })
                    continue

                if not normalized_pptx or not Path(normalized_pptx).exists():
                    logging.error(
                        f"[YD-PREP] {file_name}: .pptx не создан"
                    )
                    prepared.append({
                        "file_name": file_name,
                        "file_slug": safe_folder_name(file_name),
                        "failed_at_stage": "normalize",
                    })
                    continue
            else:
                normalized_pptx = local_pptx

            # Проверяем отмену после нормализации
            task_sess = sessions.get(task_id)
            if task_sess is None or task_sess.get("cancelled"):
                logging.info(
                    f"[YD-PREP] Задача {task_id} отменена после нормализации "
                    f"{file_name} — прерываем"
                )
                await _yd_cleanup_task(
                    task_id, session_key, task_dir,
                    owner_user_id, chat_id, nonce,
                )
                return

            # ✅ Считаем total_slides. Если python-pptx не открывает файл —
            # нормализуем его через LibreOffice в валидный .pptx, чтобы
            # и подсчёт, и последующий make_dark_mode работали.
            total_slides = None
            try:
                from pptx import Presentation
                prs = Presentation(str(normalized_pptx))
                total_slides = len(prs.slides._sldIdLst)
            except Exception as e:
                logging.warning(
                    f"[YD-PREP] {file_name}: python-pptx не открыл: {e}"
                )

            if not total_slides:
                renorm_result = await _yd_to_thread(
                    task_id,
                    librenormalize_to_pptx, normalized_pptx, per_file_dir,
                )
                if renorm_result is None:
                    logging.error(
                        f"[YD-PREP] {file_name}: файл не читается "
                        f"ни python-pptx, ни после нормализации LibreOffice"
                    )
                    prepared.append({
                        "file_name": file_name,
                        "file_slug": safe_folder_name(file_name),
                        "failed_at_stage": "count",
                    })
                    continue
                normalized_pptx, total_slides = renorm_result
                logging.info(
                    f"[YD-PREP] {file_name}: нормализован через LibreOffice "
                    f"({total_slides} слайдов)"
                )

            if not total_slides:
                logging.error(
                    f"[YD-PREP] {file_name}: total_slides == 0, пропускаем"
                )
                prepared.append({
                    "file_name": file_name,
                    "file_slug": safe_folder_name(file_name),
                    "failed_at_stage": "count",
                })
                continue

            # ✅ Читаем заметки из нормализованного .pptx
            await _safe_edit(
                status_msg,
                f"🔍 Читаю заметки докладчика: <code>{file_name_esc}</code>...",
                parse_mode="HTML",
                reply_markup=_get_cancel_keyboard(task_id).as_markup(),
            )

            notes_ok, notes, incomplete = await _yd_to_thread(
                task_id, extract_speaker_notes, str(normalized_pptx)
            )

            # Определяем проповедь по ключевым словам из конфига
            if not notes_ok:
                start, end, matches = None, None, []
                incomplete_warning = (
                    "⚠️ Не удалось прочитать заметки докладчика."
                )
            elif incomplete:
                start, end, matches = None, None, []
                incomplete_warning = (
                    "⚠️ Заметки прочитаны частично, "
                    "проповедь не определена автоматически."
                )
            else:
                keywords = yandex_state.config.sermon_keywords or [
                    yandex_state.config.sermon_keyword
                ]
                start, end, matches = find_sermon_range(notes, keywords)
                if matches and len(matches) == 1:
                    start, end = None, None
                incomplete_warning = None

            logging.info(
                f"[YD-PREP] {file_name}: заметок={len(notes)}, "
                f"matches={len(matches)}, start={start}, end={end}, "
                f"notes_ok={notes_ok}, incomplete={incomplete}"
            )

            item_ranges = [(start, end)] if start is not None else None

            prepared.append({
                "file_name": file_name,
                "file_slug": safe_folder_name(file_name),
                "file_path": normalized_pptx,
                "total_slides": total_slides,
                "start": start,
                "end": end,
                "ranges": item_ranges,
                "matches": matches,
                "notes_ok": notes_ok,
                "incomplete": incomplete,
                "incomplete_warning": incomplete_warning,
                "confirmed": False,
                "convert_mode": None,
            })

        # Публикуем pending
        cleanup_needed = False
        async with yd_session_lock:
            picker = sessions.get(session_key)
            if picker is None or picker.get("cancelled"):
                cleanup_needed = True
            else:
                task_sess = sessions.get(task_id)
                if task_sess is None or task_sess.get("cancelled"):
                    cleanup_needed = True
                else:
                    task_sess["pending"] = {
                        "task_dir": task_dir,
                        "target_base": target_base,
                        "quality": quality,
                        "prepared": prepared,
                        "owner_user_id": owner_user_id,
                        "chat_id": chat_id,
                        "session_key": session_key,
                        "nonce": nonce,
                        "bot": bot,
                        "prompt_nonce": None,
                        "prompt_idx": None,
                        "prompt_message_id": None,
                        "prompt_timeout_task": None,
                        "prompt_watchdog_nonce": None,
                        "status_message_id": status_message_id,
                    }

        if cleanup_needed:
            logging.info(f"[YD-PREP] Задача {task_id} отменена до сохранения pending")
            await _yd_cleanup_task(
                task_id, session_key, task_dir,
                owner_user_id, chat_id, nonce,
            )
            return

        needs_confirm = [
            p for p in prepared
            if "file_path" in p and not p.get("confirmed")
        ]

        logging.info(
            f"[YD-PREP] Готово: prepared={len(prepared)}, "
            f"needs_confirm={len(needs_confirm)}"
        )

        if not needs_confirm:
            for item in prepared:
                if item.get("convert_mode") is None:
                    item["convert_mode"] = "both"
            await _yd_convert_and_upload(
                bot=bot,
                task_id=task_id,
                status_msg=status_msg,
            )
            return

        await _yd_ask_sermon_confirmation(
            callback=callback,
            task_id=task_id,
            status_msg=status_msg,
            needs_confirm=needs_confirm,
        )

    except Exception as e:
        logging.error(f"[YD-PREP] Ошибка: {e}", exc_info=True)
        await _yd_cleanup_task(
            task_id, session_key, task_dir,
            owner_user_id, chat_id, nonce,
            bot=bot, status_msg=status_msg, error=e,
        )


async def _yd_ask_sermon_confirmation(
    callback: types.CallbackQuery,
    task_id: str,
    status_msg,
    needs_confirm: list,
):
    if not needs_confirm:
        return
    await _yd_render_sermon_prompt(task_id, needs_confirm[0], status_msg)


# ==========================================
# ПРОМПТЫ
# ==========================================

async def _yd_claim_prompt(callback: types.CallbackQuery) -> Optional[dict]:
    """Атомарно проверяет и 'потребляет' промпт."""
    parts = callback.data.split(":")
    if len(parts) < 4:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return None
    task_id = parts[1]
    try:
        idx = int(parts[2])
    except ValueError:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return None
    nonce = parts[3]

    session = sessions.get(task_id)
    if not session or "pending" not in session:
        await callback.answer("❌ Сессия неактивна.", show_alert=True)
        return None
    if session.get("cancelled"):
        await callback.answer("❌ Задача отменена.", show_alert=True)
        return None

    pending = session["pending"]
    if not isinstance(pending, dict):
        await callback.answer("❌ Сессия неактивна.", show_alert=True)
        return None

    if pending.get("prompt_nonce") != nonce or pending.get("prompt_idx") != idx:
        await callback.answer("⏳ Промпт уже обработан.", show_alert=True)
        return None

    if callback.from_user.id != pending["owner_user_id"]:
        await callback.answer("❌ Только автор.", show_alert=True)
        return None

    timeout_task = pending.get("prompt_timeout_task")
    if timeout_task is not None and not timeout_task.done():
        timeout_task.cancel()
    pending["prompt_timeout_task"] = None
    pending["prompt_watchdog_nonce"] = None

    pending["prompt_nonce"] = None
    pending["prompt_idx"] = None
    pending["prompt_message_id"] = None

    try:
        await callback.message.edit_reply_markup(reply_markup=None)
    except Exception:
        pass

    return pending


async def _yd_render_sermon_prompt(
    task_id: str,
    item: dict,
    status_msg,
    reply_fn=None,
) -> None:
    """
    Единая точка отрисовки промпта подтверждения проповеди.

    Три варианта:
      1. has_valid_range = True → 3 кнопки режимов + Изменить + Отмена
      2. matches > 0, но одна пометка → 2 кнопки + Изменить + Отмена
      3. matches = 0 / notes не прочитаны → «Конвертировать всё» + Изменить + Отмена
    """
    session = sessions.get(task_id)
    if not session or "pending" not in session:
        return
    if session.get("cancelled"):
        return
    pending = session["pending"]
    if not isinstance(pending, dict):
        return

    try:
        idx = pending["prepared"].index(item)
    except (ValueError, KeyError):
        logging.error(
            f"_yd_render_sermon_prompt: item не найден в prepared "
            f"для task_id={task_id}"
        )
        return

    pending["prompt_idx"] = idx

    if pending.get("prompt_nonce") is None:
        pending["prompt_nonce"] = secrets.token_hex(4)
    prompt_nonce = pending["prompt_nonce"]

    matches = item.get("matches", []) or []
    start = item.get("start")
    end = item.get("end")
    ranges = item.get("ranges")
    file_name = item.get("file_name", "")
    file_esc = html_module.escape(file_name)

    # ✅ total_slides из кэша _yd_prepare_files.
    total_slides = item.get("total_slides", 0)
    if total_slides == 0:
        file_path = item.get("file_path")
        if file_path and Path(file_path).exists():
            try:
                from pptx import Presentation
                prs = Presentation(str(file_path))
                total_slides = len(prs.slides._sldIdLst)
            except Exception as e:
                logging.warning(f"Не удалось определить число слайдов: {e}")

    kb = InlineKeyboardBuilder()

    has_valid_range = (
        bool(matches)
        and start is not None
        and end is not None
        and start <= end
    )

    logging.debug(
        f"[YD-PROMPT] task_id={task_id}, idx={idx}, file={file_name!r}, "
        f"matches={len(matches)}, start={start}, end={end}, ranges={ranges}, "
        f"has_valid_range={has_valid_range}, total_slides={total_slides}"
    )

    if has_valid_range:
        if ranges:
            sermon_count = _count_slides_in_ranges(ranges, total_slides)
            ranges_text = _format_ranges_text(ranges, start, end)
        else:
            sermon_count = _count_slides_for_range(start, end, total_slides)
            ranges_text = _format_ranges_text(None, start, end)

        other_count = max(0, total_slides - sermon_count)

        manual_range = item.get("manual_range", False)

        if manual_range:
            text = (
                f"🎯 <b>Диапазон проповеди установлен</b>\n\n"
                f"📄 Файл: <code>{file_esc}</code>\n"
                f"📊 Диапазон: <b>{ranges_text}</b>\n"
                f"🎯 Слайдов проповеди: <b>{sermon_count}</b>\n\n"
                f"❓ <b>Какие слайды конвертировать?</b>"
            )
        else:
            preview = ", ".join(str(n) for n in matches[:15])
            if len(matches) > 15:
                preview += f" …и ещё {len(matches) - 15}"

            text = (
                f"🎯 <b>Найдена пометка «проповедь»</b>\n\n"
                f"📄 Файл: <code>{file_esc}</code>\n"
                f"📌 Слайды с пометкой: <code>{preview}</code>\n"
                f"📊 Предлагаемый диапазон: <b>{ranges_text}</b>\n\n"
                f"❓ <b>Какие слайды конвертировать?</b>"
            )

        kb.row(
            InlineKeyboardButton(
                text=f"✅ Только проповедь ({sermon_count} с.)",
                callback_data=(
                    f"yd_sermon_mode:{task_id}:{idx}:{prompt_nonce}:sermon"
                ),
            ),
        )
        kb.row(
            InlineKeyboardButton(
                text=f"📄 Только остальные ({other_count} с.)",
                callback_data=(
                    f"yd_sermon_mode:{task_id}:{idx}:{prompt_nonce}:other"
                ),
            ),
        )
        kb.row(
            InlineKeyboardButton(
                text=f"📦 Проповеди и остальное раздельно ({total_slides} с.)",
                callback_data=(
                    f"yd_sermon_mode:{task_id}:{idx}:{prompt_nonce}:both"
                ),
            ),
        )
        kb.row(
            InlineKeyboardButton(
                text="✏️ Изменить диапазон Проповеди",
                callback_data=f"yd_sermon_edit:{task_id}:{idx}:{prompt_nonce}",
            ),
        )
    else:
        notes_ok = item.get("notes_ok", True)
        incomplete = item.get("incomplete", False)

        if not notes_ok:
            reason = (
                "⚠️ <b>Не удалось прочитать заметки докладчика.</b>\n"
                "Возможно, файл повреждён или содержит только изображения."
            )
        elif incomplete:
            reason = (
                "⚠️ <b>Заметки прочитаны частично.</b>\n"
                "Автоматически определить проповедь не удалось."
            )
        elif matches:
            match_str = ", ".join(str(n) for n in matches[:5])
            reason = (
                f"📌 Найдена <b>одна</b> пометка «проповедь»: "
                f"слайд <code>{match_str}</code>\n"
                f"Для одной пометки авто-диапазон не строится."
            )
        else:
            reason = (
                "В заметках докладчика нет слова «проповедь»."
            )

        text = (
            f"🤔 <b>Автоматически определить проповедь не удалось</b>\n\n"
            f"📄 Файл: <code>{file_esc}</code>\n"
            f"📌 Всего слайдов: <b>{total_slides}</b>\n\n"
            f"{reason}\n\n"
            f"❓ <b>Что делать с файлом?</b>"
        )

        kb.row(
            InlineKeyboardButton(
                text=f"📄 Конвертировать всё ({total_slides} с.)",
                callback_data=(
                    f"yd_sermon_mode:{task_id}:{idx}:{prompt_nonce}:other"
                ),
            ),
        )
        kb.row(
            InlineKeyboardButton(
                text="✏️ Указать диапазон Проповеди",
                callback_data=f"yd_sermon_edit:{task_id}:{idx}:{prompt_nonce}",
            ),
        )

    kb.row(
        InlineKeyboardButton(
            text="❌ Отменить задачу",
            callback_data=f"yd_task_cancel:{task_id}",
        ),
    )

    sent_msg = None
    if reply_fn is not None:
        sent_msg = await reply_fn(text, parse_mode="HTML", reply_markup=kb.as_markup())
    elif status_msg is not None:
        await status_msg.edit_text(text, parse_mode="HTML", reply_markup=kb.as_markup())
        sent_msg = status_msg
    else:
        logging.warning(
            f"_yd_render_sermon_prompt: нет ни reply_fn, ни status_msg для {task_id}"
        )
        return

    if pending.get("prompt_nonce") != prompt_nonce:
        logging.info(
            f"_yd_render_sermon_prompt: nonce изменён конкурентно для {task_id}"
        )
        return

    if sent_msg is not None and hasattr(sent_msg, "message_id"):
        pending["prompt_message_id"] = sent_msg.message_id

    existing_task = pending.get("prompt_timeout_task")
    existing_nonce = pending.get("prompt_watchdog_nonce")

    if existing_task is not None and not existing_task.done():
        if existing_nonce == prompt_nonce:
            return
        existing_task.cancel()
        pending["prompt_timeout_task"] = None
        pending["prompt_watchdog_nonce"] = None

    pending["prompt_timeout_task"] = asyncio.create_task(
        _yd_prompt_timeout_watchdog(
            task_id, yandex_state.config.prompt_timeout_sec, prompt_nonce
        )
    )
    pending["prompt_watchdog_nonce"] = prompt_nonce


async def _yd_prompt_timeout_watchdog(task_id: str, timeout_sec: int, expected_nonce: str):
    """Если пользователь не ответил на промпт — уведомляем и очищаем."""
    try:
        await asyncio.sleep(timeout_sec)
        session = sessions.get(task_id)
        if not session or "pending" not in session:
            return
        pending = session["pending"]
        if not isinstance(pending, dict):
            return
        if pending.get("prompt_nonce") != expected_nonce:
            return

        chat_id = pending.get("chat_id")
        bot: Optional[Bot] = pending.get("bot")
        owner_user_id = pending.get("owner_user_id")
        session_key = pending.get("session_key")
        task_dir = pending.get("task_dir")
        nonce = pending.get("nonce")
        prompt_message_id = pending.get("prompt_message_id")

        logging.info(
            f"[YD-PROMPT] ⏰ Промпт {task_id} не подтверждён за {timeout_sec}s — очистка"
        )

        if bot is not None and chat_id is not None:
            minutes = max(1, timeout_sec // 60)
            try:
                await bot.send_message(
                    chat_id=chat_id,
                    text=(
                        f"⏰ <b>Время ожидания истекло</b> ({minutes} мин).\n\n"
                        f"Задача отменена, временные файлы удалены.\n"
                        f"Если хотите обработать файл — запустите /sunday заново."
                    ),
                    parse_mode="HTML",
                )
            except Exception as e:
                logging.warning(f"Не удалось уведомить о таймауте: {e}")

        if bot is not None and chat_id is not None and prompt_message_id is not None:
            try:
                await bot.edit_message_reply_markup(
                    chat_id=chat_id,
                    message_id=prompt_message_id,
                    reply_markup=None,
                )
            except Exception:
                pass

        await _yd_cleanup_task(
            task_id, session_key, task_dir,
            owner_user_id, chat_id, nonce,
        )
    except asyncio.CancelledError:
        pass
    except Exception as e:
        logging.error(f"Ошибка watchdog {task_id}: {e}", exc_info=True)


# ==========================================
# ОЧИСТКА
# ==========================================

async def _yd_cleanup_task(
    task_id: str,
    session_key: str,
    task_dir: Path,
    owner_user_id: int,
    chat_id: int,
    nonce: str,
    bot: Optional[Bot] = None,
    status_msg=None,
    error: Optional[Exception] = None,
):
    """Идемпотентная очистка."""
    logging.info(
        f"[YD-CLEANUP] task_id={task_id}, session_key={session_key!r}, "
        f"reason={'error' if error else 'normal'}"
    )

    # ✅ Гарантируем, что спиннер мёртв до редактирования сообщения
    try:
        await _yd_stop_spinner(task_id)
    except Exception as e:
        logging.debug(f"[YD-CLEANUP] _yd_stop_spinner: {e}")

    # Снимаем клавиатуру со всех сообщений задачи
    try:
        session = sessions.get(task_id)
        if session is not None:
            pending = session.get("pending")
            if isinstance(pending, dict):
                task_bot = pending.get("bot")
                task_chat_id = pending.get("chat_id")
                if task_bot is not None and task_chat_id is not None:
                    message_ids = []
                    for key in ("prompt_message_id", "status_message_id"):
                        mid = pending.get(key)
                        if mid is not None and mid not in message_ids:
                            message_ids.append(mid)

                    for mid in message_ids:
                        try:
                            await task_bot.edit_message_reply_markup(
                                chat_id=task_chat_id,
                                message_id=mid,
                                reply_markup=None,
                            )
                            logging.debug(
                                f"[YD-CLEANUP] Снята клавиатура с сообщения {mid}"
                            )
                        except Exception as e:
                            if "message is not modified" in str(e):
                                logging.debug(
                                    f"[YD-CLEANUP] {mid}: клавиатура уже снята"
                                )
                            else:
                                logging.debug(
                                    f"[YD-CLEANUP] Не удалось снять клавиатуру "
                                    f"с {mid}: {e}"
                                )
    except Exception as e:
        logging.debug(f"[YD-CLEANUP] Ошибка снятия клавиатуры: {e}")

    # Папка задачи
    in_flight_ops = _active_ops.get(task_id)
    if in_flight_ops:
        ops_snapshot = list(in_flight_ops)
        logging.info(
            f"[YD-CLEANUP] {len(ops_snapshot)} in-flight ops для {task_id} — "
            f"откладываем удаление task_dir"
        )
        asyncio.create_task(
            _yd_deferred_task_dir_cleanup(task_dir, ops_snapshot)
        )
    else:
        try:
            if task_dir and task_dir.exists():
                shutil.rmtree(task_dir)
                logging.info(f"🧹 Удалена папка задачи: {task_dir}")
        except Exception as e:
            logging.error(f"Ошибка удаления task_dir {task_dir}: {e}")

    # Watchdog
    try:
        session = sessions.get(task_id)
        if session is not None:
            pending = session.get("pending")
            if isinstance(pending, dict):
                timeout_task = pending.get("prompt_timeout_task")
                current_task = asyncio.current_task()
                if (
                    timeout_task is not None
                    and not timeout_task.done()
                    and timeout_task is not current_task
                ):
                    timeout_task.cancel()
                pending["prompt_timeout_task"] = None
                pending["prompt_watchdog_nonce"] = None
    except Exception as e:
        logging.error(f"Ошибка отмены watchdog для {task_id}: {e}", exc_info=True)

    # Сессии
    try:
        async with yd_session_lock:
            picker = sessions.get(session_key)
            # ✅ Мутируем picker ТОЛЬКО если он относится к нашей задаче
            # (nonce совпадает). Иначе это picker новой сессии, и его
            # processing=True отражает реально идущий захват — сброс
            # флага откроет гонку дубликатов.
            if picker is not None and picker.get("nonce") == nonce:
                picker["processing"] = False
                task_ids = picker.get("task_ids")
                if isinstance(task_ids, list) and task_id in task_ids:
                    task_ids.remove(task_id)
                sessions.pop(session_key, None)
            elif picker is not None:
                # Picker чужой. Если наш task_id как-то в него попал —
                # убираем, но processing не трогаем.
                task_ids = picker.get("task_ids")
                if isinstance(task_ids, list) and task_id in task_ids:
                    task_ids.remove(task_id)
                    logging.warning(
                        f"[YD-CLEANUP] task_id={task_id} оказался в чужом "
                        f"picker'е (nonce={picker.get('nonce')!r}, "
                        f"ожидался {nonce!r}) — удаляем"
                    )

            sessions.pop(task_id, None)
            yd_active_tasks.discard(task_id)
    except Exception as e:
        logging.error(f"Ошибка очистки сессий для {task_id}: {e}", exc_info=True)

    # yd_release
    try:
        await yd_release(owner_user_id, chat_id, nonce)
    except Exception as e:
        logging.error(f"Ошибка yd_release для {task_id}: {e}", exc_info=True)

    # Сообщение об ошибке
    if error is not None and bot is not None and status_msg is not None:
        try:
            await status_msg.edit_text(
                f"❌ <b>Ошибка обработки</b>\n\n"
                f"<code>{html_module.escape(str(error)[:200])}</code>\n\n"
                f"Временные файлы удалены. Попробуйте снова.",
                parse_mode="HTML",
            )
        except Exception as e:
            logging.error(f"Ошибка отправки сообщения для {task_id}: {e}")


# ==========================================
# КОНВЕРТАЦИЯ + UPLOAD (после выбора режима)
# ==========================================

async def _yd_convert_and_upload(
    bot: Bot,
    task_id: str,
    status_msg,
):
    """
    Вторая фаза: конвертация PNG + упаковка в ZIP + upload.
    Вызывается после выбора режима в yd_sermon_mode.
    """
    session = sessions.get(task_id)
    if not session or "pending" not in session:
        return

    def _is_cancelled() -> bool:
        s = sessions.get(task_id)
        return s is None or s.get("cancelled", False)

    if _is_cancelled():
        logging.info(f"[YD-UP] Задача {task_id} отменена — upload пропущен")
        pending = session.get("pending")
        if isinstance(pending, dict):
            await _yd_cleanup_task(
                task_id,
                pending.get("session_key"),
                pending.get("task_dir"),
                pending.get("owner_user_id"),
                pending.get("chat_id"),
                pending.get("nonce"),
            )
        return

    pending = session["pending"]
    if not isinstance(pending, dict):
        return

    # ✅ Регистрируем worker task для отмены из yd_task_cancel_callback.
    # Если пользователь отменит во время upload_file, cancel handler
    # отменит и дождётся этой задачи ПЕРЕД yd_release — иначе in-flight
    # upload (overwrite=True) может завершиться уже ПОСЛЕ того, как
    # перезапущенная задача загрузит свой файл на тот же remote path,
    # и перезапишет его устаревшим контентом.
    worker_task = asyncio.current_task()
    session["worker_task"] = worker_task

    prepared = pending["prepared"]
    target_base = pending["target_base"]
    task_dir = pending["task_dir"]
    owner_user_id = pending["owner_user_id"]
    chat_id = pending["chat_id"]
    session_key = pending["session_key"]
    nonce = pending["nonce"]
    quality = pending.get("quality", "2k")

    if status_msg is not None and hasattr(status_msg, "message_id"):
        pending["status_message_id"] = status_msg.message_id

    zip_tmp_dir: Optional[Path] = None
    cleanup_done = False

    logging.info(
        f"[YD-UP] Старт: task_id={task_id}, файлов={len(prepared)}, "
        f"target_base={target_base!r}"
    )

    try:
        zip_tmp_dir = Path(tempfile.mkdtemp(prefix=f"pptx2png_{task_id}_"))
        logging.debug(f"[YD-UP] Создана временная папка {zip_tmp_dir}")

        total_uploaded_zip = 0
        total_slides_packed = 0
        total_failed = 0
        report_lines = [f"📁 Обработано файлов: <b>{len(prepared)}</b>\n"]
        links_by_folder: dict[str, str] = {}

        for f_idx, item in enumerate(prepared, start=1):
            _touch_task(task_dir)

            if _is_cancelled():
                logging.info(f"[YD-UP] Отмена перед файлом #{f_idx}")
                return

            file_name = item["file_name"]
            file_name_esc = html_module.escape(file_name)
            file_slug = item["file_slug"]

            if item.get("convert_mode") == "skip":
                logging.info(f"[YD-UP] {file_name}: пропущен пользователем")
                report_lines.append(f"⏭ {file_name_esc} — пропущен пользователем")
                continue

            if item.get("failed_at_stage"):
                stage = item["failed_at_stage"]
                stage_text = {
                    "download":  "ошибка скачивания",
                    "normalize": "ошибка подготовки .ppt → .pptx",
                    "count":     "не удалось определить число слайдов",
                    "convert":   "ошибка конвертации",
                    "no_pngs":   "нет PNG",
                }.get(stage, stage)
                report_lines.append(f"❌ {file_name_esc} — {stage_text}")
                total_failed += 1
                continue

            file_path = item.get("file_path")
            if not file_path or not Path(file_path).exists():
                logging.error(f"[YD-UP] {file_name}: PPTX не найден ({file_path})")
                report_lines.append(f"❌ {file_name_esc} — файл PPTX не найден")
                total_failed += 1
                continue

            start = item.get("start")
            end = item.get("end")
            ranges = item.get("ranges")
            convert_mode = item.get("convert_mode", "both")
            incomplete_warning = item.get("incomplete_warning")

            logging.info(
                f"[YD-UP] Файл #{f_idx}: {file_name!r}, "
                f"convert_mode={convert_mode}, ranges={ranges}"
            )

            mode_label = _mode_label(convert_mode)
            ranges_str = _format_ranges_text(ranges, start, end)

            header_lines = [
                f"🎬 <b>{mode_label}</b>",
                f"📄 Файл: <code>{file_name_esc}</code>",
            ]
            if ranges or (start is not None and end is not None):
                header_lines.append(f"📊 Диапазон: <code>{ranges_str}</code>")
            header_lines.append("")
            header = "\n".join(header_lines)

            base_convert_text = f"{header}\n⚙️ Конвертирую в PNG"

            await _safe_edit(
                status_msg,
                base_convert_text,
                parse_mode="HTML",
                reply_markup=_get_cancel_keyboard(task_id).as_markup(),
            )

            temp_png_dir = task_dir / f"png_{f_idx}"
            temp_png_dir.mkdir(exist_ok=True)

            try:
                pngs, used_pptx = await _yd_with_spinner(
                    status_msg,
                    task_id,
                    base_convert_text,
                    _yd_run_protected(
                        task_id,
                        convert_all_pngs(
                            Path(file_path), temp_png_dir, quality
                        ),
                    ),
                )
            except Exception as e:
                logging.error(
                    f"[YD-UP] {file_name}: ошибка конвертации: {e}",
                    exc_info=True,
                )
                report_lines.append(f"❌ {file_name_esc} — ошибка конвертации")
                total_failed += 1
                continue

            if not pngs:
                logging.warning(f"[YD-UP] {file_name}: не создано ни одного PNG")
                report_lines.append(f"❌ {file_name_esc} — нет PNG")
                total_failed += 1
                continue

            pngs_sorted = sorted(pngs, key=lambda p: p.name)

            if _is_cancelled():
                logging.info(f"[YD-UP] Отмена после конвертации {file_name}")
                return

            sermon_pngs = []
            other_pngs = []
            for slide_idx, png_path in enumerate(pngs_sorted, start=1):
                if _is_sermon_slide(item, slide_idx):
                    sermon_pngs.append(png_path)
                else:
                    other_pngs.append(png_path)

            ranges_text = ranges_str

            pptx2png_dir = (
                f"{target_base}/{yandex_state.config.pptx2png_folder}/{file_slug}"
            )
            sermon_dir = f"{target_base}/{yandex_state.config.sermon_folder}"

            need_pptx2png_folder = convert_mode in ("other", "both")
            need_sermon_folder = convert_mode in ("sermon", "both")

            if need_pptx2png_folder:
                ok1 = await yandex_state.config.client.ensure_folder(pptx2png_dir)
                if not ok1:
                    logging.error(
                        f"[YD-UP] {file_name}: не удалось создать {pptx2png_dir!r}"
                    )
                    report_lines.append(
                        f"❌ {file_name_esc} — не удалось создать папки"
                    )
                    total_failed += 1
                    continue

            if need_sermon_folder and sermon_pngs:
                ok2 = await yandex_state.config.client.ensure_folder(sermon_dir)
                if not ok2:
                    logging.error(
                        f"[YD-UP] {file_name}: не удалось создать {sermon_dir!r}"
                    )
                    report_lines.append(
                        f"❌ {file_name_esc} — не удалось создать папку проповеди"
                    )
                    total_failed += 1
                    continue

            # === Архив с проповедью ===
            sermon_info = None
            if need_sermon_folder and sermon_pngs:
                if _is_cancelled():
                    logging.info(f"[YD-UP] Отмена перед ZIP проповеди")
                    return

                sermon_zip_name = f"{file_slug}_проповедь.zip"
                sermon_zip_path = zip_tmp_dir / sermon_zip_name

                try:
                    await _yd_to_thread(
                        task_id,
                        create_zip_stream, sermon_pngs, sermon_zip_path,
                    )
                except Exception as e:
                    logging.error(
                        f"[YD-UP] {file_name}: ошибка ZIP проповеди: {e}",
                        exc_info=True,
                    )
                    report_lines.append(
                        f"❌ {file_name_esc} — ошибка упаковки проповеди"
                    )
                    total_failed += 1
                    continue

                if _is_cancelled():
                    _safe_unlink(sermon_zip_path)
                    return

                sermon_zip_size = sermon_zip_path.stat().st_size
                remote_path = f"{sermon_dir}/{sermon_zip_name}"

                base_upload_text = (
                    f"{header}\n"
                    f"📤 Загружаю архив проповеди "
                    f"(<code>{file_name_esc}</code>)"
                )

                await _safe_edit(
                    status_msg,
                    base_upload_text,
                    parse_mode="HTML",
                    reply_markup=_get_cancel_keyboard(task_id).as_markup(),
                )

                if _is_cancelled():
                    _safe_unlink(sermon_zip_path)
                    return

                ok = await _yd_with_spinner(
                    status_msg,
                    task_id,
                    base_upload_text,
                    yandex_state.config.client.upload_file(
                        sermon_zip_path, remote_path
                    ),
                )

                if ok:
                    total_uploaded_zip += 1
                    total_slides_packed += len(sermon_pngs)
                    links_by_folder[sermon_dir] = "🎯 Проповедь"
                    sermon_info = (
                        f"🎯 Проповедь ({ranges_text}): {len(sermon_pngs)} слайдов → "
                        f"<code>{html_module.escape(yandex_state.config.sermon_folder)}/"
                        f"{html_module.escape(sermon_zip_name)}</code> "
                        f"({_format_size(sermon_zip_size)})"
                    )
                    logging.info(
                        f"[YD-UP] {file_name}: sermon ZIP загружен "
                        f"({len(sermon_pngs)} слайдов)"
                    )
                else:
                    total_failed += 1
                    sermon_info = (
                        f"❌ Проповедь ({ranges_text}): "
                        f"не удалось загрузить ZIP"
                    )
                    logging.error(f"[YD-UP] {file_name}: sermon ZIP failed")

                _safe_unlink(sermon_zip_path)
                for png in sermon_pngs:
                    _safe_unlink(png)

            # === Архив с остальными слайдами ===
            other_info = None
            if need_pptx2png_folder and other_pngs:
                if _is_cancelled():
                    logging.info(f"[YD-UP] Отмена перед ZIP слайдов")
                    return

                other_zip_name = f"{file_slug}_слайды.zip"
                other_zip_path = zip_tmp_dir / other_zip_name

                try:
                    await _yd_to_thread(
                        task_id,
                        create_zip_stream, other_pngs, other_zip_path,
                    )
                except Exception as e:
                    logging.error(
                        f"[YD-UP] {file_name}: ошибка ZIP слайдов: {e}",
                        exc_info=True,
                    )
                    report_lines.append(
                        f"❌ {file_name_esc} — ошибка упаковки слайдов"
                    )
                    total_failed += 1
                    if sermon_info:
                        report_lines.append(f"{f_idx}. 📄 <b>{file_name_esc}</b>")
                        report_lines.append(f"   • {sermon_info}")
                    continue

                if _is_cancelled():
                    _safe_unlink(other_zip_path)
                    return

                other_zip_size = other_zip_path.stat().st_size
                remote_path = f"{pptx2png_dir}/{other_zip_name}"

                base_upload_text = (
                    f"{header}\n"
                    f"📤 Загружаю архив слайдов "
                    f"(<code>{file_name_esc}</code>)"
                )

                await _safe_edit(
                    status_msg,
                    base_upload_text,
                    parse_mode="HTML",
                    reply_markup=_get_cancel_keyboard(task_id).as_markup(),
                )

                if _is_cancelled():
                    _safe_unlink(other_zip_path)
                    return

                ok = await _yd_with_spinner(
                    status_msg,
                    task_id,
                    base_upload_text,
                    yandex_state.config.client.upload_file(
                        other_zip_path, remote_path
                    ),
                )

                other_folder_short = (
                    f"{yandex_state.config.pptx2png_folder}/{file_slug}"
                )

                if ok:
                    total_uploaded_zip += 1
                    total_slides_packed += len(other_pngs)
                    links_by_folder[pptx2png_dir] = "📄 Остальные слайды"
                    other_info = (
                        f"📄 Остальные: {len(other_pngs)} слайдов → "
                        f"<code>{html_module.escape(other_folder_short)}/"
                        f"{html_module.escape(other_zip_name)}</code> "
                        f"({_format_size(other_zip_size)})"
                    )
                    logging.info(
                        f"[YD-UP] {file_name}: slides ZIP загружен "
                        f"({len(other_pngs)} слайдов)"
                    )
                else:
                    total_failed += 1
                    other_info = (
                        "❌ Остальные слайды: не удалось загрузить ZIP"
                    )
                    logging.error(f"[YD-UP] {file_name}: slides ZIP failed")

                _safe_unlink(other_zip_path)
                for png in other_pngs:
                    _safe_unlink(png)

            if not sermon_info and not other_info:
                if convert_mode == "sermon" and not sermon_pngs:
                    other_info = "ℹ️ Проповедь не найдена в этом файле"
                elif convert_mode == "other" and not other_pngs:
                    other_info = "ℹ️ Все слайды относятся к проповеди"
                else:
                    other_info = "ℹ️ Нет слайдов для конвертации"

            entry_lines = [f"{f_idx}. 📄 <b>{file_name_esc}</b>"]
            if sermon_info:
                entry_lines.append(f"   • {sermon_info}")
            if other_info:
                entry_lines.append(f"   • {other_info}")
            if incomplete_warning:
                entry_lines.append(f"   • {incomplete_warning}")
            report_lines.append("\n".join(entry_lines))

        if _is_cancelled():
            logging.info(f"[YD-UP] Отмена перед отправкой отчёта")
            return

        if total_failed > 0:
            report_lines.append(
                f"\n⚠️ Всего загружено архивов: <b>{total_uploaded_zip}</b>\n"
                f"📊 Всего слайдов: <b>{total_slides_packed}</b>\n"
                f"❌ Ошибок: <b>{total_failed}</b>"
            )
        else:
            report_lines.append(
                f"\n📊 Всего загружено архивов: <b>{total_uploaded_zip}</b>\n"
                f"📊 Всего слайдов: <b>{total_slides_packed}</b>"
            )

        if links_by_folder:
            report_lines.append("\n🔗 <b>Ссылки на Яндекс.Диск:</b>")
            for folder_path, label in links_by_folder.items():
                url = _yd_public_url(folder_path)
                folder_name = folder_path.rsplit("/", 1)[-1]
                report_lines.append(
                    f'   • {label}: <a href="{url}">'
                    f'{html_module.escape(folder_name)}/</a>'
                )

        try:
            await status_msg.edit_reply_markup(reply_markup=None)
        except Exception:
            pass

        await _yd_send_report(
            bot=bot,
            chat_id=chat_id,
            status_msg=status_msg,
            report_lines=report_lines,
            total_uploaded=total_uploaded_zip,
            total_failed=total_failed,
        )

        logging.info(
            f"[YD-UP] Итог: zip={total_uploaded_zip}, "
            f"slides={total_slides_packed}, failed={total_failed}"
        )

    except Exception as e:
        logging.error(f"[YD-UP] Ошибка: {e}", exc_info=True)
        cleanup_done = True
        await _yd_cleanup_task(
            task_id, session_key, task_dir,
            owner_user_id, chat_id, nonce,
            bot=bot, status_msg=status_msg, error=e,
        )
        return
    finally:
        # ✅ Снимаем регистрацию worker'а — чтобы cancel handler не
        # пытался отменить уже завершённый таск.
        try:
            sess = sessions.get(task_id)
            if sess is not None and sess.get("worker_task") is worker_task:
                sess["worker_task"] = None
        except Exception:
            pass

        if zip_tmp_dir is not None and zip_tmp_dir.exists():
            try:
                shutil.rmtree(zip_tmp_dir)
                logging.debug(f"[YD-UP] Удалена временная папка {zip_tmp_dir}")
            except Exception as e:
                logging.warning(
                    f"[YD-UP] Не удалось удалить {zip_tmp_dir}: {e}",
                    exc_info=True,
                )

        if not cleanup_done:
            await _yd_cleanup_task(
                task_id, session_key, task_dir,
                owner_user_id, chat_id, nonce,
            )


async def _yd_send_report(
    bot: Bot,
    chat_id: int,
    status_msg,
    report_lines: list,
    total_uploaded: int,
    total_failed: int,
):
    MAX_MSG_LEN = 3500
    chunks = []
    current_chunk = []
    current_len = 0

    for line in report_lines:
        line_len = len(line) + 1
        if current_len + line_len > MAX_MSG_LEN and current_chunk:
            chunks.append("\n".join(current_chunk))
            current_chunk = [line]
            current_len = line_len
        else:
            current_chunk.append(line)
            current_len += line_len

    if current_chunk:
        chunks.append("\n".join(current_chunk))

    delivered = 0
    failed_chunks = []
    first_delivered = False

    if chunks:
        try:
            await status_msg.edit_text(
                chunks[0], parse_mode="HTML", disable_web_page_preview=True
            )
            delivered += 1
            first_delivered = True
        except Exception as e:
            logging.error(f"Ошибка edit_text первой части: {e}", exc_info=True)
            try:
                await bot.send_message(
                    chat_id=chat_id,
                    text=chunks[0],
                    parse_mode="HTML",
                    disable_web_page_preview=True,
                )
                delivered += 1
                first_delivered = True
            except Exception as e2:
                logging.error(f"Не удалось отправить первую часть: {e2}", exc_info=True)
                failed_chunks.append(1)

    if not first_delivered:
        try:
            fallback = f"⚠️ Не удалось показать полный отчёт.\n📊 Загружено: {total_uploaded}"
            if total_failed:
                fallback += f"\n❌ Ошибок: {total_failed}"
            await bot.send_message(chat_id=chat_id, text=fallback)
        except Exception as e:
            logging.error(f"Не удалось отправить fallback-отчёт: {e}", exc_info=True)

    for i, chunk in enumerate(chunks[1:], start=2):
        try:
            await bot.send_message(
                chat_id=chat_id,
                text=chunk,
                parse_mode="HTML",
                disable_web_page_preview=True,
            )
            delivered += 1
        except Exception as e:
            logging.error(f"Ошибка отправки части {i}: {e}", exc_info=True)
            failed_chunks.append(i)

    if failed_chunks and delivered > 0:
        try:
            await bot.send_message(
                chat_id=chat_id,
                text=(
                    f"⚠️ Не удалось доставить {len(failed_chunks)} "
                    f"из {len(chunks)} частей отчёта. Проверьте Яндекс.Диск."
                ),
            )
        except Exception as e:
            logging.error(f"Не удалось отправить предупреждение: {e}")


# ==========================================
# ХЕНДЛЕРЫ
# ==========================================

@router.message(Command("sunday"))
async def cmd_sunday(message: types.Message, check_access, bot: Bot):
    if not await check_access(message):
        return

    if yandex_state.config.client is None:
        await message.reply("❌ Яндекс.Диск не настроен. Обратитесь к администратору.")
        return

    if not yandex_state.config.base_path:
        await message.reply("❌ Не задан base_path Яндекс.Диска в settings.ini.")
        return

    logging.info(
        f"[YD] /sunday от user={message.from_user.id}, "
        f"base_path={yandex_state.config.base_path!r}"
    )

    # ✅ Не позволяем плодить параллельные задачи. Используем общую
    # семантику через _picker_is_active — отменённые задачи не блокируют.
    session_key = f"yd_{message.from_user.id}_{message.chat.id}"
    async with yd_session_lock:
        existing_picker = sessions.get(session_key)
        is_active = _picker_is_active(existing_picker)

    if is_active:
        await message.reply(
            "⚠️ <b>У вас уже есть активная задача.</b>\n\n"
            "Дождитесь её завершения или отмените командой /cancel_yd, "
            "затем запустите /sunday снова.",
            parse_mode="HTML",
        )
        return

    nonce = await yd_try_acquire(message.from_user.id, message.chat.id)
    if nonce is None:
        await message.reply(
            "⏳ У вас уже активна сессия подготовки трансляции.\n"
            "Дождитесь завершения или нажмите /cancel_yd."
        )
        return

    status_msg = None
    session_created = False

    try:
        status_msg = await message.reply("🔍 Проверяю Яндекс.Диск...")

        ok, err = await yandex_state.config.client.check_access()
        if not ok:
            await status_msg.edit_text(
                f"❌ <b>Яндекс.Диск недоступен</b>\n\n"
                f"Причина: <code>{html_module.escape(str(err))}</code>\n\n"
                f"Проверьте токен в <code>config.ini</code>.",
                parse_mode="HTML",
            )
            return

        sunday = get_nearest_sunday()
        sunday_str = sunday.strftime("%d.%m.%Y")
        month_str = month_folder_name(sunday)

        await status_msg.edit_text(
            f"✅ Яндекс.Диск доступен\n"
            f"📅 Ближайшее воскресенье: <b>{html_module.escape(sunday_str)}</b>\n"
            f"📁 Ожидаемая папка: "
            f"<code>{html_module.escape(f'{month_str}/{sunday_str}')}</code>\n\n"
            f"🔍 Проверяю структуру папок...",
            parse_mode="HTML",
        )

        paths = await resolve_sunday_paths(
            yandex_state.config.client,
            yandex_state.config.base_path,
            sunday,
            yandex_state.config.source_folder,
            yandex_state.config.target_folder,
        )
        if not paths:
            src_esc = html_module.escape(yandex_state.config.source_folder)
            tgt_esc = html_module.escape(yandex_state.config.target_folder)
            await status_msg.edit_text(
                f"❌ <b>Структура папок не найдена</b>\n\n"
                f"Ожидалось:\n"
                f"<code>{html_module.escape(yandex_state.config.base_path)}/</code>\n"
                f"<code>  {html_module.escape(month_str)}/</code>\n"
                f"<code>    {html_module.escape(sunday_str)}/</code>\n"
                f"<code>      {src_esc}/</code>\n"
                f"<code>      {tgt_esc}/</code>",
                parse_mode="HTML",
            )
            return

        try:
            pptx_files = await find_pptx_in_source(
                yandex_state.config.client, paths["source"], sunday
            )
        except YandexDiskError as e:
            logging.error(f"[YD] Ошибка доступа к источнику: {e}")
            await status_msg.edit_text(
                f"❌ <b>Ошибка обращения к Яндекс.Диску</b>\n\n"
                f"<code>{html_module.escape(str(e))}</code>\n\n"
                f"Попробуйте позже.",
                parse_mode="HTML",
            )
            return

        logging.info(f"[YD] Найдено pptx: {len(pptx_files)}")

        if not pptx_files:
            src_esc = html_module.escape(yandex_state.config.source_folder)
            await status_msg.edit_text(
                f"📅 Ближайшее воскресенье: "
                f"<b>{html_module.escape(sunday_str)}</b>\n"
                f"📍 Папка: <code>{html_module.escape(paths['source'])}</code>\n\n"
                f"❌ <b>pptx-файлы не найдены.</b>\n\n"
                f"Положите pptx с датой <code>{sunday:%d.%m.%y}</code> "
                f"в папку <code>{src_esc}</code> и попробуйте снова.",
                parse_mode="HTML",
            )
            return

        MAX_LEN = 3500
        header_lines = [
            f"📅 Ближайшее воскресенье: <b>{html_module.escape(sunday_str)}</b>",
            f"📍 Папка: <code>{html_module.escape(paths['source'])}</code>",
            "",
            f"📄 <b>Найдено файлов: {len(pptx_files)}</b>",
            "",
        ]
        body_lines = []
        omitted = 0
        for idx, f in enumerate(pptx_files, start=1):
            size_mb = f.get("size", 0) / (1024 * 1024)
            name_escaped = html_module.escape(f["name"])
            line = f"{idx}. <code>{name_escaped}</code> — {size_mb:.1f} МБ"
            candidate = "\n".join(header_lines + body_lines + [line])
            if len(candidate) > MAX_LEN:
                omitted = len(pptx_files) - idx + 1
                break
            body_lines.append(line)

        if omitted > 0:
            body_lines.append("")
            body_lines.append(f"…и ещё <b>{omitted}</b> файл(ов) не показано.")

        body_lines.append("")
        body_lines.append("🎬 Выберите файл для обработки:")

        if not await yd_is_active(message.from_user.id, message.chat.id, nonce):
            logging.info(
                f"[YD] Сессия {message.from_user.id}:{message.chat.id} "
                f"была отменена во время выполнения"
            )
            try:
                await status_msg.edit_text("❌ Операция отменена пользователем.")
            except Exception:
                pass
            return

        session_key = f"yd_{message.from_user.id}_{message.chat.id}"

        race_detected = False
        async with yd_session_lock:
            # ✅ Защитный путь: сюда не должны попасть из-за проверки
            # в начале cmd_sunday. Если попали — это гонка, отказываемся
            # перезаписывать активную сессию. Используем _picker_is_active
            # для единой семантики (отменённые задачи игнорируются).
            old_picker = sessions.get(session_key)
            if _picker_is_active(old_picker):
                logging.warning(
                    f"[YD] Race: active picker {session_key!r} "
                    f"processing={old_picker.get('processing')}, "
                    f"task_ids={old_picker.get('task_ids')!r}"
                )
                race_detected = True
            elif old_picker is not None:
                sessions.pop(session_key, None)

            if not race_detected:
                sessions[session_key] = {
                    "user_id": message.from_user.id,
                    "chat_id": message.chat.id,
                    "sunday": sunday,
                    "sunday_str": sunday_str,
                    "paths": paths,
                    "files": pptx_files,
                    "nonce": nonce,
                    "created_at": time.time(),
                    "task_ids": [],
                    "cancelled": False,
                    "processing": False,
                }

        if race_detected:
            try:
                await status_msg.edit_text(
                    "⚠️ <b>Другая команда /sunday уже активна.</b>\n\n"
                    "Дождитесь её завершения или отмените /cancel_yd.",
                    parse_mode="HTML",
                )
            except Exception:
                await message.reply(
                    "⚠️ Другая команда /sunday уже активна."
                )
            return

        kb = InlineKeyboardBuilder()
        for idx, f in enumerate(pptx_files):
            prefix = "🎯" if "служение" in f["name"].lower() else "📄"
            kb.row(InlineKeyboardButton(
                text=f"{prefix} {f['name']}",
                callback_data=f"yd_pick:{message.from_user.id}:{nonce}:{idx}",
            ))
        if len(pptx_files) > 1:
            kb.row(InlineKeyboardButton(
                text="📁 Все подряд",
                callback_data=f"yd_pick:{message.from_user.id}:{nonce}:all",
            ))
        kb.row(InlineKeyboardButton(
            text="❌ Отмена",
            callback_data=f"yd_cancel:{message.from_user.id}:{nonce}",
        ))

        await status_msg.edit_text(
            "\n".join(header_lines + body_lines),
            parse_mode="HTML",
            reply_markup=kb.as_markup(),
        )
        session_created = True

        await yd_release(message.from_user.id, message.chat.id, nonce)
        logging.info(
            f"[YD] 🔓 Сессия {message.from_user.id}:{message.chat.id} "
            f"освобождена после показа списка (nonce={nonce[:8]}...)"
        )

    except Exception as e:
        logging.error(f"[YD] Ошибка cmd_sunday: {e}", exc_info=True)
        try:
            if status_msg:
                await status_msg.edit_text(
                    f"❌ Ошибка: <code>{html_module.escape(str(e)[:200])}</code>",
                    parse_mode="HTML",
                )
            else:
                await message.reply(f"❌ Ошибка: {str(e)[:200]}")
        except Exception:
            pass
    finally:
        if not session_created:
            released = await yd_release(message.from_user.id, message.chat.id, nonce)
            if released:
                logging.info(
                    f"[YD] 🔓 Сессия {message.from_user.id}:{message.chat.id} "
                    f"освобождена (неудачный запуск, nonce={nonce[:8]}...)"
                )


@router.callback_query(F.data.startswith("yd_pick:"))
async def yd_pick(
    callback: types.CallbackQuery,
    bot: Bot,
    SHM_DIR: str,
    user_mgr,
):
    parts = callback.data.split(":")
    if len(parts) != 4:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return

    try:
        owner_user_id = int(parts[1])
    except ValueError:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return

    callback_nonce = parts[2]
    file_selector = parts[3]

    if callback.from_user.id != owner_user_id:
        await callback.answer("❌ Только автор запроса может выбрать файл.", show_alert=True)
        return

    session_key = f"yd_{owner_user_id}_{callback.message.chat.id}"

    async with yd_session_lock:
        session = sessions.get(session_key)
        if not session or session.get("nonce") != callback_nonce:
            await callback.answer("❌ Сессия неактивна.", show_alert=True)
            return

        # ✅ Единая семантика через _picker_is_active.
        if _picker_is_active(session):
            await callback.answer("⏳ Обработка уже запущена.", show_alert=True)
            return

        session["processing"] = True
        files = session["files"]
        sunday = session["sunday"]
        paths = session["paths"]

        if file_selector == "all":
            files_to_process = files
        else:
            try:
                idx = int(file_selector)
                if idx < 0 or idx >= len(files):
                    session.pop("processing", None)
                    await callback.answer("❌ Файл не найден.", show_alert=True)
                    return
                files_to_process = [files[idx]]
            except ValueError:
                session.pop("processing", None)
                await callback.answer("❌ Некорректный выбор.", show_alert=True)
                return

    # ✅ try/except вокруг запуска подготовки: если callback.answer()
    # или _yd_prepare_files упадут до регистрации task_id — сбрасываем
    # processing, иначе следующий /sunday навсегда упрётся в отказ.
    try:
        if file_selector == "all":
            await callback.answer("⏳ Обрабатываю все файлы...")
        else:
            await callback.answer("⏳ Начинаю обработку...")

        await _yd_prepare_files(
            callback=callback,
            bot=bot,
            SHM_DIR=SHM_DIR,
            user_mgr=user_mgr,
            files_to_process=files_to_process,
            sunday=sunday,
            paths=paths,
            session_key=session_key,
            nonce=callback_nonce,
        )
    except Exception as e:
        logging.error(
            f"[YD-PICK] ошибка запуска подготовки: {e}",
            exc_info=True,
        )
        try:
            async with yd_session_lock:
                picker = sessions.get(session_key)
                if (
                    picker is not None
                    and picker.get("nonce") == callback_nonce
                ):
                    picker["processing"] = False
        except Exception as cleanup_err:
            logging.error(
                f"[YD-PICK] не удалось сбросить processing: {cleanup_err}",
                exc_info=True,
            )
        try:
            await callback.message.reply(
                f"❌ Не удалось запустить обработку:\n"
                f"<code>{html_module.escape(str(e)[:200])}</code>\n\n"
                f"Запустите <code>/sunday</code> заново.",
                parse_mode="HTML",
            )
        except Exception:
            pass


@router.callback_query(F.data.startswith("yd_cancel:"))
async def yd_cancel_callback(callback: types.CallbackQuery):
    parts = callback.data.split(":")
    if len(parts) != 3:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return

    try:
        owner_user_id = int(parts[1])
    except ValueError:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return

    callback_nonce = parts[2]

    if callback.from_user.id != owner_user_id:
        await callback.answer(
            "❌ Только автор запроса может отменить операцию.", show_alert=True
        )
        return

    session_key = f"yd_{owner_user_id}_{callback.message.chat.id}"

    is_stale = False
    task_ids_to_cancel = []
    async with yd_session_lock:
        session = sessions.get(session_key)
        if session is None or session.get("nonce") != callback_nonce:
            is_stale = True
        else:
            session["cancelled"] = True
            task_ids_to_cancel = list(session.get("task_ids", []))
            sessions.pop(session_key, None)

    if is_stale:
        await yd_release(owner_user_id, callback.message.chat.id, callback_nonce)
        try:
            await callback.message.edit_text("❌ Сессия уже неактивна.")
        except Exception:
            pass
        await callback.answer("❌ Сессия уже неактивна.", show_alert=True)
        return

    for tid in task_ids_to_cancel:
        async with yd_session_lock:
            task_sess = sessions.get(tid)
            if task_sess is not None:
                task_sess["cancelled"] = True

    await yd_release(owner_user_id, callback.message.chat.id, callback_nonce)

    try:
        await callback.message.edit_text("❌ Операция отменена.")
    except Exception:
        pass
    await callback.answer()


@router.message(Command("cancel_yd"))
async def cmd_cancel_yd(message: types.Message, check_access):
    if not await check_access(message):
        return

    session_key = f"yd_{message.from_user.id}_{message.chat.id}"

    task_ids_to_cancel = []
    session = None
    async with yd_session_lock:
        session = sessions.get(session_key)
        if session is not None:
            session["cancelled"] = True
            task_ids_to_cancel = list(session.get("task_ids", []))
            sessions.pop(session_key, None)

    for tid in task_ids_to_cancel:
        async with yd_session_lock:
            task_sess = sessions.get(tid)
            if task_sess is not None:
                task_sess["cancelled"] = True

    released = await yd_release(message.from_user.id, message.chat.id)

    if session is None and not task_ids_to_cancel and released:
        await message.reply("ℹ️ У вас нет активной сессии Яндекс.Диска.")
        return

    await message.reply("✅ Сессия Яндекс.Диска сброшена.")


# ==========================================
# ВЫБОР РЕЖИМА КОНВЕРТАЦИИ
# ==========================================

@router.callback_query(F.data.startswith("yd_sermon_mode:"))
async def yd_sermon_mode(callback: types.CallbackQuery, bot: Bot):
    """
    Пользователь выбрал режим конвертации:
      - sermon → только проповедь
      - other  → только остальные
      - both   → оба ZIP раздельно
    """
    parts = callback.data.split(":")
    if len(parts) != 5:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return

    task_id = parts[1]
    try:
        idx = int(parts[2])
    except ValueError:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return
    nonce = parts[3]
    mode = parts[4]

    if mode not in ("sermon", "other", "both"):
        await callback.answer("❌ Неизвестный режим.", show_alert=True)
        return

    session = sessions.get(task_id)
    if not session or "pending" not in session:
        await callback.answer("❌ Сессия неактивна.", show_alert=True)
        return

    claimed = await _yd_claim_prompt(callback)
    if claimed is None:
        return
    pending = claimed

    item = pending["prepared"][idx]

    if mode in ("sermon", "both") and (
        item.get("start") is None or item.get("end") is None
    ):
        await callback.answer(
            "❌ Диапазон не задан. Укажите его вручную.",
            show_alert=True,
        )
        await _yd_render_sermon_prompt(task_id, item, callback.message)
        return

    item["convert_mode"] = mode
    item["confirmed"] = True

    try:
        await callback.answer("⏳ Принято, начинаю конвертацию...")
    except Exception:
        pass

    remaining = [
        p for p in pending["prepared"]
        if "file_path" in p and not p.get("confirmed")
    ]

    if remaining:
        await _yd_render_sermon_prompt(task_id, remaining[0], callback.message)
        return

    await _yd_convert_and_upload(
        bot=bot,
        task_id=task_id,
        status_msg=callback.message,
    )


# ==========================================
# ОТМЕНА ТЕКУЩЕЙ ЗАДАЧИ
# ==========================================

@router.callback_query(F.data.startswith("yd_task_cancel:"))
async def yd_task_cancel_callback(callback: types.CallbackQuery):
    """
    Отмена текущей Yandex-задачи.

    Работает на любой стадии:
      - До старта (pending=None) → ставим cancelled, задача отменится
        на ближайшей проверке.
      - На промпте → cleanup сразу.
      - На стадии upload → текущий шаг доводится до конца, потом cleanup.
    """
    parts = callback.data.split(":")
    if len(parts) != 2:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return

    task_id = parts[1]

    session = sessions.get(task_id)
    if not session:
        logging.info(
            f"[YD-TASK-CANCEL] Задача {task_id!r} не найдена (уже очищена)"
        )
        try:
            await callback.message.edit_reply_markup(reply_markup=None)
        except Exception:
            pass

        await callback.answer(
            "ℹ️ Эта задача уже завершена или отменена.",
            show_alert=True,
        )
        return

    owner_user_id = session.get("user_id")
    if callback.from_user.id != owner_user_id:
        await callback.answer(
            "❌ Только автор задачи может её отменить.",
            show_alert=True,
        )
        return

    # ✅ Ставим cancelled=True ВСЕГДА
    session["cancelled"] = True

    # ✅ Немедленно убираем task_id из picker.task_ids — иначе он
    # будет блокировать новый /sunday, пока cleanup не закончится
    # (а cleanup ждёт окончания convert_all_pngs, что может быть
    # десятки секунд).
    session_key_for_task = session.get("session_key")
    if session_key_for_task:
        async with yd_session_lock:
            picker = sessions.get(session_key_for_task)
            if picker is not None:
                task_ids = picker.get("task_ids")
                if isinstance(task_ids, list) and task_id in task_ids:
                    task_ids.remove(task_id)

    # ✅ Сначала отменяем и дожидаемся worker'а. Иначе in-flight
    # upload_file (overwrite=True) может финишировать уже ПОСЛЕ того,
    # как перезапущенная задача загрузит свой архив на тот же
    # remote path, и перезапишет его устаревшим контентом.
    worker = session.get("worker_task")
    current_task = asyncio.current_task()
    if worker is not None and not worker.done() and worker is not current_task:
        logging.info(
            f"[YD-TASK-CANCEL] Отменяем worker {task_id} перед yd_release"
        )
        try:
            worker.cancel()
        except Exception as e:
            logging.debug(f"[YD-TASK-CANCEL] worker.cancel: {e}")
        try:
            await asyncio.wait_for(worker, timeout=15.0)
            logging.info(
                f"[YD-TASK-CANCEL] worker {task_id} завершён"
            )
        except asyncio.TimeoutError:
            logging.warning(
                f"[YD-TASK-CANCEL] worker {task_id} не завершился за 15s — "
                f"продолжаем, но in-flight upload может продолжиться в фоне"
            )
        except asyncio.CancelledError:
            pass
        except Exception as e:
            logging.debug(f"[YD-TASK-CANCEL] await worker: {e}")

    # ✅ Освобождаем низкоуровневый session-lock (yd_active_sessions).
    # Делаем это ТОЛЬКО после того, как worker завершён — иначе
    # следующий /sunday может запустить upload параллельно со старым.
    # yd_release внутри сверяет nonce — если сессию успели заменить,
    # чужой nonce он не тронет.
    owner_chat_id = session.get("chat_id")
    owner_nonce = session.get("nonce")
    if owner_chat_id is not None and owner_nonce:
        try:
            released = await yd_release(
                owner_user_id, owner_chat_id, owner_nonce
            )
            logging.info(
                f"[YD-TASK-CANCEL] yd_release для {task_id}: "
                f"released={released}"
            )
        except Exception as e:
            logging.error(
                f"[YD-TASK-CANCEL] yd_release упал для {task_id}: {e}",
                exc_info=True,
            )

    # ✅ ГЛАВНОЕ: останавливаем активный спиннер ДО любых edit_text,
    # иначе следующий tick перезапишет сообщение об отмене.
    await _yd_stop_spinner(task_id)

    pending = session.get("pending")

    # Случай 1: pending ещё не создан (подготовка не завершена)
    if not isinstance(pending, dict):
        logging.info(
            f"[YD-TASK-CANCEL] Задача {task_id!r} отменена на стадии подготовки"
        )

        try:
            await callback.answer("❌ Отмена запрошена")
        except Exception:
            pass

        try:
            await callback.message.edit_text(
                "⏳ <b>Отмена запрошена…</b>\n\n"
                "Задача будет отменена после текущего шага.\n"
                "<i>Больше ничего нажимать не нужно.</i>",
                parse_mode="HTML",
                reply_markup=None,
            )
        except Exception:
            pass

        return

    # Случай 2: pending есть
    logging.info(
        f"[YD-TASK-CANCEL] Пользователь {callback.from_user.id} "
        f"отменил задачу {task_id}"
    )

    try:
        await callback.answer("❌ Задача отменена")
    except Exception:
        pass

    timeout_task = pending.get("prompt_timeout_task")
    if timeout_task is not None and not timeout_task.done():
        timeout_task.cancel()
    pending["prompt_timeout_task"] = None
    pending["prompt_watchdog_nonce"] = None

    prompt_active = pending.get("prompt_nonce") is not None

    if prompt_active:
        try:
            await callback.message.edit_text(
                "❌ <b>Задача отменена</b>\n\n"
                "Временные файлы удалены.\n"
                "Запустите <code>/sunday</code> заново, если нужно.",
                parse_mode="HTML",
                reply_markup=None,
            )
        except Exception:
            pass

        await _yd_cleanup_task(
            task_id=task_id,
            session_key=pending.get("session_key"),
            task_dir=pending.get("task_dir"),
            owner_user_id=owner_user_id,
            chat_id=pending.get("chat_id"),
            nonce=pending.get("nonce"),
        )
    else:
        try:
            await callback.message.edit_text(
                "⏳ <b>Отмена запрошена…</b>\n\n"
                "Текущий шаг завершится, затем задача будет очищена.\n"
                "<i>Больше ничего нажимать не нужно.</i>",
                parse_mode="HTML",
                reply_markup=None,
            )
        except Exception:
            pass


# ==========================================
# ИЗМЕНЕНИЕ ДИАПАЗОНА ПРОПОВЕДИ
# ==========================================

@router.callback_query(F.data.startswith("yd_sermon_edit:"))
async def yd_sermon_edit(callback: types.CallbackQuery, bot: Bot):
    parts = callback.data.split(":")
    if len(parts) != 4:
        await callback.answer("❌ Некорректный запрос.", show_alert=True)
        return
    task_id = parts[1]

    session = sessions.get(task_id)
    if not session or "pending" not in session:
        await callback.answer("❌ Сессия неактивна.", show_alert=True)
        return

    claimed = await _yd_claim_prompt(callback)
    if claimed is None:
        return
    pending = claimed

    idx = int(parts[2])
    manual_nonce = secrets.token_hex(4)

    current = pending["prepared"][idx]
    file_name_esc = html_module.escape(current["file_name"])

    total_slides = current.get("total_slides", 0)
    if total_slides == 0:
        file_path = current.get("file_path")
        if file_path and Path(file_path).exists():
            try:
                from pptx import Presentation
                prs = Presentation(str(file_path))
                total_slides = len(prs.slides._sldIdLst)
            except Exception as e:
                logging.warning(f"Не удалось определить число слайдов: {e}")

    matches = current.get("matches", []) or []
    start = current.get("start")
    end = current.get("end")
    ranges = current.get("ranges")
    notes_ok = current.get("notes_ok", True)
    incomplete = current.get("incomplete", False)

    context_lines = []
    if matches:
        preview = ", ".join(str(n) for n in matches[:20])
        if len(matches) > 20:
            preview += f" …и ещё {len(matches) - 20}"
        context_lines.append(
            f"📌 <b>Найдены пометки на слайдах:</b> <code>{preview}</code>"
        )
        if start is not None and end is not None:
            ranges_text = _format_ranges_text(ranges, start, end)
            context_lines.append(
                f"📊 <b>Предложенный диапазон:</b> <code>{ranges_text}</code>"
            )
    elif not notes_ok:
        context_lines.append("⚠️ <i>Заметки докладчика не удалось прочитать.</i>")
    elif incomplete:
        context_lines.append("⚠️ <i>Заметки прочитаны частично.</i>")
    else:
        context_lines.append(
            "📌 <i>Автоматических пометок «проповедь» не найдено.</i>"
        )

    context_block = "\n".join(context_lines)
    if context_block:
        context_block = "\n\n" + context_block

    cancel_kb = InlineKeyboardBuilder()
    cancel_kb.row(
        InlineKeyboardButton(
            text="❌ Отменить задачу",
            callback_data=f"yd_task_cancel:{task_id}",
        ),
    )

    sent_msg = None
    try:
        await callback.answer()
        sent_msg = await callback.message.edit_text(
            f"✏️ <b>Введите диапазон проповеди</b>\n\n"
            f"📄 Файл: <code>{file_name_esc}</code>\n"
            f"📊 Всего слайдов: <b>{total_slides}</b>"
            f"{context_block}\n\n"
            f"<b>Формат:</b> <code>5-30</code> или <code>5,7,10-15</code>\n"
            f"Отправьте текстом в чат (ответом на это сообщение).\n"
            f"<i>Отправьте <code>отмена</code> или <code>0</code>, чтобы пропустить файл.</i>",
            parse_mode="HTML",
            reply_markup=cancel_kb.as_markup(),
        )
    except Exception as e:
        logging.error(
            f"yd_sermon_edit: edit_text упал для {task_id}: {e}",
            exc_info=True,
        )
        await _yd_cleanup_task(
            task_id,
            pending.get("session_key"),
            pending.get("task_dir"),
            pending.get("owner_user_id"),
            pending.get("chat_id"),
            pending.get("nonce"),
            bot=bot,
            status_msg=None,
            error=None,
        )
        try:
            await bot.send_message(
                chat_id=callback.message.chat.id,
                text=(
                    "❌ Не удалось показать форму ввода. Задача сброшена."
                ),
            )
        except Exception:
            pass
        return

    pending["awaiting_range_for_idx"] = idx
    pending["prompt_nonce"] = manual_nonce
    pending["prompt_idx"] = idx
    if sent_msg is not None and hasattr(sent_msg, "message_id"):
        pending["prompt_message_id"] = sent_msg.message_id

    if pending.get("prompt_nonce") != manual_nonce:
        logging.info(
            f"yd_sermon_edit: nonce изменён конкурентно для {task_id}"
        )
        return

    old_timeout = pending.get("prompt_timeout_task")
    if old_timeout is not None and not old_timeout.done():
        old_timeout.cancel()

    pending["prompt_timeout_task"] = asyncio.create_task(
        _yd_prompt_timeout_watchdog(
            task_id, yandex_state.config.prompt_timeout_sec, manual_nonce
        )
    )
    pending["prompt_watchdog_nonce"] = manual_nonce


# ==========================================
# ПУБЛИЧНЫЕ ОБЁРТКИ ДЛЯ handlers.py
# ==========================================

render_sermon_prompt = _yd_render_sermon_prompt
convert_and_upload = _yd_convert_and_upload
cleanup_task = _yd_cleanup_task
is_sermon_slide = _is_sermon_slide
claim_prompt = _yd_claim_prompt
prompt_timeout_watchdog = _yd_prompt_timeout_watchdog